"""Build the Workshop Readout pptx template the corpus is missing.

    uv run python scripts/build_workshop_readout_template.py

The must-have covers Word, Excel *and* PowerPoint, but YMSLI supplied no pptx
template -- only a hackathon introduction deck, which is a presentation, not a
form to generate from. Plan §1.5 calls for building one, and this script is it
rather than a committed binary: a generated artifact stays reviewable, and the
deck's content comes from the real corpus so the demo doesn't look synthetic.

The deck is deliberately shaped to exercise every region kind the detector has
to handle, and every part kind the compiler has to preserve:

  slide 1  title + subtitle + the YMSLI logo lifted from the MOM  -> title, field, branding
  slide 2  agenda bullet list                                     -> repeating_block
  slides 3-5  decision slides, title + bullets                    -> repeating_block
  slide 6  action items table, header row + one empty row         -> table_loop
  slide 7  a chart, which brings its own embedded workbook part    -> derived

The empty table row on slide 6 is the point, not an oversight: it is the row
template `clone_table_rows` copies, and it makes the table a `table_loop` (a
header with nothing under it) exactly like MOM row 27.
"""

import io
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches, Pt

FILES_DIR = Path(__file__).resolve().parent.parent.parent / "Files"
MOM_PATH = FILES_DIR / "MOM Sample.xlsx"
OUTPUT_PATH = FILES_DIR / "Workshop Readout Template.pptx"

# Layout indices in python-pptx's default Office theme.
TITLE_LAYOUT = 0
TITLE_AND_CONTENT_LAYOUT = 1
TITLE_ONLY_LAYOUT = 5

ACTION_COLUMNS = ["Sl.No.", "Action", "Responsibility", "Target date", "Status"]

AGENDA = [
    "Scope and objectives for this cycle",
    "Current-state process walkthrough",
    "System demonstration",
    "Decisions taken",
    "Open points and next steps",
]

DECISIONS = [
    (
        "Decision: Invoice uniqueness key",
        [
            "Vendor plus invoice number is treated as the uniqueness key.",
            "Rationale: matches the existing manual control and needs no master data change.",
            "Impact: duplicate detection can be switched on without a data cleanse.",
        ],
    ),
    (
        "Decision: Non-PO invoice routing",
        [
            "Non-PO invoices route through a separate approval workflow.",
            "Rationale: cost centre coding happens at line level and needs a different approver set.",
            "Impact: one additional workflow to configure and test.",
        ],
    ),
    (
        "Decision: Payment terms default",
        [
            "Net 45 is the default, overridable per vendor.",
            "Rationale: reflects the majority of existing vendor agreements.",
            "Impact: vendor master requires a terms review before go-live.",
        ],
    ),
]

# Fit/Gap counts by module, taken from the BFL sample's own Fit/Gap column.
FIT_GAP = {
    "General Ledger": (8, 1),
    "Bank Management": (3, 0),
    "Budget Management": (1, 1),
    "Fixed Assets": (5, 2),
    "Accounts Payable": (4, 3),
}


def _logo_bytes() -> bytes | None:
    """The YMSLI logo, reused from the MOM so the deck carries real branding."""
    if not MOM_PATH.exists():
        return None
    with zipfile.ZipFile(MOM_PATH) as archive:
        if "xl/media/image1.png" not in archive.namelist():
            return None
        return archive.read("xl/media/image1.png")


def _set_bullets(placeholder, items: list[str]) -> None:
    text_frame = placeholder.text_frame
    text_frame.text = items[0]
    for item in items[1:]:
        paragraph = text_frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
    for paragraph in text_frame.paragraphs:
        paragraph.font.size = Pt(18)


def build() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)  # 16:9, what the corporate deck uses
    presentation.slide_height = Inches(7.5)

    # -- slide 1: title -------------------------------------------------
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_LAYOUT])
    slide.shapes.title.text = "Workshop Readout"
    subtitle = slide.placeholders[1]
    subtitle.text_frame.text = (
        "Fixed Asset - Acquire to Dispose  |  YMSLI / YMESG  |  Date to be filled"
    )
    # Formatting set on the *run* rather than the paragraph, deliberately. The
    # rest of the deck styles paragraphs, so without this nothing would test
    # whether the compiler preserves run-level `<a:rPr>` when it overwrites
    # existing text -- which is the one thing python-pptx's own text setter
    # throws away.
    subtitle_run = subtitle.text_frame.paragraphs[0].runs[0]
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.italic = True
    logo = _logo_bytes()
    if logo is not None:
        slide.shapes.add_picture(
            io.BytesIO(logo),
            left=Emu(presentation.slide_width - Inches(2.1)),
            top=Inches(0.35),
            width=Inches(1.6),
        )
    slide.notes_slide.notes_text_frame.text = (
        "Readout deck generated from the meeting minutes. Confirm attendees before sharing."
    )

    # -- slide 2: agenda ------------------------------------------------
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[TITLE_AND_CONTENT_LAYOUT]
    )
    slide.shapes.title.text = "Agenda"
    _set_bullets(slide.placeholders[1], AGENDA)

    # -- slides 3-5: decisions -----------------------------------------
    for title, points in DECISIONS:
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[TITLE_AND_CONTENT_LAYOUT]
        )
        slide.shapes.title.text = title
        _set_bullets(slide.placeholders[1], points)

    # -- slide 6: action items table -----------------------------------
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_ONLY_LAYOUT])
    slide.shapes.title.text = "Action Items"
    table_shape = slide.shapes.add_table(
        rows=2,  # header plus one empty row: the row template to clone
        cols=len(ACTION_COLUMNS),
        left=Inches(0.7),
        top=Inches(1.8),
        width=presentation.slide_width - Inches(1.4),
        height=Inches(1.0),
    )
    table = table_shape.table
    for index, heading in enumerate(ACTION_COLUMNS):
        cell = table.cell(0, index)
        cell.text = heading
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
    for index in range(len(ACTION_COLUMNS)):
        table.cell(1, index).text_frame.paragraphs[0].font.size = Pt(14)
    table.columns[0].width = Inches(0.9)
    table.columns[1].width = Inches(5.2)
    table.columns[2].width = Inches(2.4)
    table.columns[3].width = Inches(1.9)
    table.columns[4].width = Inches(1.5)

    # -- slide 7: chart -------------------------------------------------
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_ONLY_LAYOUT])
    slide.shapes.title.text = "Fit / Gap by Module"
    chart_data = CategoryChartData()
    chart_data.categories = list(FIT_GAP)
    chart_data.add_series("Fit", tuple(fit for fit, _ in FIT_GAP.values()))
    chart_data.add_series("Gap", tuple(gap for _, gap in FIT_GAP.values()))
    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        x=Inches(1.0),
        y=Inches(1.8),
        cx=presentation.slide_width - Inches(2.0),
        cy=Inches(4.8),
        chart_data=chart_data,
    )
    chart = chart_frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    slide.notes_slide.notes_text_frame.text = (
        "Counts come from the Business Function List Fit/Gap column."
    )

    return presentation


def main() -> int:
    presentation = build()
    presentation.save(OUTPUT_PATH)

    with zipfile.ZipFile(OUTPUT_PATH) as archive:
        parts = archive.namelist()
    print(f"wrote {OUTPUT_PATH} ({OUTPUT_PATH.stat().st_size:,} bytes, {len(parts)} parts)")
    print(f"  slides:     {sum(1 for p in parts if p.startswith('ppt/slides/slide'))}")
    print(f"  media:      {[p for p in parts if '/media/' in p]}")
    print(f"  charts:     {[p for p in parts if p.startswith('ppt/charts/')]}")
    print(f"  embeddings: {[p for p in parts if p.startswith('ppt/embeddings/')]}")
    print(f"  notes:      {sum(1 for p in parts if p.startswith('ppt/notesSlides/notes'))}")
    if not any("/media/" in part for part in parts):
        print("  WARNING: no logo embedded -- MOM Sample.xlsx was not readable")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
