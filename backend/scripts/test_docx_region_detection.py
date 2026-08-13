"""Region-detector report for the docx sample. Read-only on the source.

Run this after touching app/ingestion/regions_docx.py:

    uv run python scripts/test_docx_region_detection.py

It prints every detected region so a human can eyeball recall, then measures
recall against a ground truth built **independently of the detector** -- the
expectations below are derived from the document with plain python-docx and raw
XML (paragraphs holding a `<w:drawing>`, runs overriding the default font size,
text beginning `Area Path:`, a table header containing `Ver.`), never from
`regions_docx`'s own output. So a miss here is a real miss, and the printed
recall figure is the plan's 85% target measured rather than asserted.

It also synthesises a small document for the two shapes the BRD has no example
of -- an empty table and a caption with nothing under it -- rather than shipping
those branches untested.
"""

import io
import re
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx  # noqa: E402
from docx.shared import Pt  # noqa: E402

from app.ingestion.fingerprint_docx import fingerprint_docx  # noqa: E402
from app.ingestion.regions_docx import (  # noqa: E402
    DocumentSchema,
    Region,
    detect_regions_docx,
)

FILES_DIR = Path(__file__).resolve().parent.parent.parent / "Files"
SAMPLE = "BRD Sample.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

RECALL_TARGET = 0.85  # plan §1.4


def _short(text: str | None, width: int = 34) -> str:
    if not text:
        return ""
    flat = " ".join(text.split())
    return flat if len(flat) <= width else flat[: width - 1] + "…"


def report(name: str, schema: DocumentSchema, elapsed_ms: float) -> None:
    print("=" * 108)
    print(f"{name}   detector={schema.detector_version}   {elapsed_ms:.0f}ms")
    print()
    for key in (
        "body_paragraphs",
        "body_size_pt",
        "heading_ladder",
        "styles_used",
        "drawings",
        "hyperlink_paragraphs",
        "export_tool",
        "house_pattern",
        "in_doc_version",
    ):
        if key in schema.notes:
            print(f"  {key:22} {schema.notes[key]}")
    print()
    for region in schema.regions:
        flags = "".join(
            ("E" if region.editable else "-", "P" if region.protected else "-")
        )
        extra = ""
        if region.kind in ("table_loop", "table_data", "history"):
            cols = [c["name"] for c in region.notes.get("columns", [])]
            extra = (
                f"  rows={region.notes.get('row_count')}"
                f" template_row={region.notes.get('template_row')}"
                f" cols={_short(', '.join(cols), 52)}"
            )
        elif region.kind == "repeating_block":
            extra = (
                f"  items={region.notes['item_count']}"
                f" numbered={region.notes['numbered']}"
                f" insert_at={region.notes['insert_at_para']}"
            )
        elif region.kind == "branding":
            extra = f"  drawings={region.notes.get('drawings')} anchor={region.notes.get('anchor')}"
        elif not region.editable and region.notes.get("reason"):
            extra = f"  reason={region.notes['reason']}"
        print(
            f"    {flags} {region.kind:<15} {region.locator.ref:<14}"
            f" {_short(region.label):<35} {_short(region.current_content, 26)}{extra}"
        )
    print()
    print(f"  totals: {schema.counts_by_kind()}")
    print(f"  writable regions: {len(schema.writable())} / {len(schema.regions)}")


# ── ground truth, built from the document rather than from the detector ────


def _expected_regions(data: bytes) -> list[dict]:
    """What a reader of this document would say is in it.

    Every rule here is a statement about the file that can be checked with
    python-docx alone. None of them consult `regions_docx`.
    """
    document = docx.Document(io.BytesIO(data))
    paragraphs = document.paragraphs
    default_pt = _default_size_pt(data)
    expected: list[dict] = []

    def size_of(paragraph) -> float:
        sizes = [
            run.font.size.pt
            for run in paragraph.runs
            if run.font.size is not None and run.text.strip()
        ]
        return max(sizes) if sizes else default_pt

    sizes_present = sorted(
        {size_of(p) for p in paragraphs if p.text.strip() and size_of(p) > default_pt},
        reverse=True,
    )

    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        # Compared unstripped against unstripped: many paragraphs in this export
        # begin with a space (`' Area Path: ...'`), and stripping one side turns
        # every one of them into a false "text is inside a hyperlink" verdict.
        runs_own_text = "".join(run.text for run in paragraph.runs) == paragraph.text

        # An image paragraph is branding, whether the drawing is inline or
        # floating. python-docx's inline_shapes can only see one of the six.
        if paragraph._p.find(f".//{W}drawing") is not None:
            expected.append(
                {"what": f"drawing in para {index}", "paragraphs": [index], "kind": "branding"}
            )
            continue

        if not text:
            continue

        # Text held inside a hyperlink cannot be rewritten at run level, so the
        # fail-safe kind is the correct answer for it.
        if not runs_own_text and not paragraph.runs:
            expected.append(
                {
                    "what": f"hyperlink-only para {index}",
                    "paragraphs": [index],
                    "kind": "immutable",
                }
            )
            continue

        size = size_of(paragraph)
        if size in sizes_present[:1]:
            expected.append(
                {"what": f"document title para {index}", "paragraphs": [index], "kind": "title"}
            )
            continue

        match = re.match(r"^([A-Za-z][A-Za-z .%/()\-]{1,40}):[ \t]*(\S.*)$", text)
        if match and runs_own_text:
            expected.append(
                {
                    "what": f"field {match.group(1)!r} para {index}",
                    "paragraphs": [index],
                    "kind": "field",
                }
            )
            continue

        if size > default_pt:
            expected.append(
                {"what": f"heading para {index}", "paragraphs": [index], "kind": "title"}
            )

    # Every contiguous numbering group is a list a generator would rewrite.
    for group in _numbering_groups(paragraphs):
        expected.append(
            {
                "what": f"list numId={group['num_id']} paras {group['start']}-{group['end']}",
                "paragraphs": list(range(group["start"], group["end"] + 1)),
                "kind": "repeating_block",
            }
        )

    # The prose under each `Description:` / `Acceptance Criteria:` caption.
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not re.match(r"^(Description|Acceptance Criteria):$", text, re.IGNORECASE):
            continue
        members = []
        for cursor in range(index + 1, len(paragraphs)):
            following = paragraphs[cursor].text.strip()
            if not following:
                if members:
                    break
                continue
            if paragraphs[cursor]._p.find(f".//{W}drawing") is not None:
                break
            if "".join(r.text for r in paragraphs[cursor].runs) != paragraphs[cursor].text:
                break
            if size_of(paragraphs[cursor]) > default_pt or following.endswith(":"):
                break
            members.append(cursor)
        if members:
            expected.append(
                {
                    "what": f"{text!r} block under para {index}",
                    "paragraphs": members,
                    "kind": "repeating_block",
                }
            )

    for table_index, table in enumerate(document.tables):
        header = [cell.text.strip() for cell in table.rows[0].cells]
        normalised = {re.sub(r"[^a-z0-9]", "", h.lower()) for h in header}
        if {"ver", "version"} & normalised:
            kind = "history"
        elif len(table.rows) > 1 and any(c.text.strip() for c in table.rows[1].cells):
            kind = "table_data"
        else:
            kind = "table_loop"
        expected.append(
            {"what": f"table {table_index} ({kind})", "table": table_index, "kind": kind}
        )

    return expected


def _default_size_pt(data: bytes) -> float:
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        styles = archive.read("word/styles.xml").decode("utf-8", "replace")
    match = re.search(
        r"<w:docDefaults>.*?<w:rPrDefault>.*?<w:sz\s+w:val=\"(\d+)\"", styles, re.DOTALL
    )
    return int(match.group(1)) / 2.0 if match else 11.0


def _numbering_groups(paragraphs: list) -> list[dict]:
    groups: list[dict] = []
    current: dict | None = None
    for index, paragraph in enumerate(paragraphs):
        properties = paragraph._p.pPr
        num_id = None
        if properties is not None and properties.numPr is not None:
            if properties.numPr.numId is not None:
                num_id = int(properties.numPr.numId.val)
        if num_id is None:
            current = None
            continue
        if current and current["num_id"] == num_id and index == current["end"] + 1:
            current["end"] = index
        else:
            current = {"num_id": num_id, "start": index, "end": index}
            groups.append(current)
    return groups


def _covers(region: Region, para_index: int) -> bool:
    """Does this region account for that paragraph?

    A region's caption counts as belonging to it. The caption is the region's
    name rather than a write target, so it is deliberately outside `para_span`
    -- rewriting a block must not overwrite the heading that introduces it --
    but it *is* classified, and a coverage check that ignored it would report
    `Description:` as an unclassified paragraph.
    """
    locator = region.locator
    if locator.para_index == para_index:
        return True
    if locator.para_span and locator.para_span[0] <= para_index <= locator.para_span[1]:
        return True
    return region.notes.get("caption_para") == para_index


def measure_recall(schema: DocumentSchema, expected: list[dict]) -> tuple[int, int, list[str]]:
    hits = 0
    misses: list[str] = []
    for item in expected:
        if "table" in item:
            found = any(
                region.locator.table_index == item["table"] and region.kind == item["kind"]
                for region in schema.regions
            )
        else:
            found = all(
                any(
                    _covers(region, para) and region.kind == item["kind"]
                    for region in schema.regions
                )
                for para in item["paragraphs"]
            )
        if found:
            hits += 1
        else:
            actual = sorted(
                {
                    region.kind
                    for region in schema.regions
                    if (
                        region.locator.table_index == item.get("table")
                        if "table" in item
                        else any(_covers(region, p) for p in item["paragraphs"])
                    )
                }
            )
            misses.append(f"{item['what']}: want {item['kind']}, detected {actual or 'nothing'}")
    return hits, len(expected), misses


# ── assertions the plan names by name ─────────────────────────────────────


def check_brd(schema: DocumentSchema, data: bytes) -> list[str]:
    failures: list[str] = []
    by_kind: dict[str, list[Region]] = {}
    for region in schema.regions:
        by_kind.setdefault(region.kind, []).append(region)

    fingerprint = fingerprint_docx(data)

    # Plan §1.2 ⑤: 6 drawings, of which inline_shapes sees 1.
    branding = by_kind.get("branding", [])
    if len(branding) != 6:
        failures.append(f"{len(branding)} branding regions, want 6 (one per <w:drawing>)")
    if sum(r.notes.get("drawings", 0) for r in branding) != fingerprint.drawings:
        failures.append("branding regions do not account for every <w:drawing>")
    inline_shapes = len(docx.Document(io.BytesIO(data)).inline_shapes)
    if inline_shapes != 1:
        failures.append(f"expected the inline_shapes trap to return 1, got {inline_shapes}")
    if any(r.editable for r in branding):
        failures.append("a branding region is editable")

    # Plan §1.3: the house revision block, and the version it declares.
    history = by_kind.get("history", [])
    if len(history) != 1:
        failures.append(f"{len(history)} history regions, want 1 (the revision block)")
    else:
        block = history[0]
        names = [c["name"] for c in block.notes["columns"]]
        if "Ver." not in names or "Updated by" not in names:
            failures.append(f"revision block columns wrong: {names}")
        if block.notes.get("in_doc_version") != "1.0":
            failures.append(
                f"in-document version read as {block.notes.get('in_doc_version')!r}, want '1.0'"
            )
        if block.notes.get("template_row") != 2:
            failures.append(
                f"revision block template_row={block.notes.get('template_row')}, want 2 "
                "(row 1 has only 5 of 9 cells; cloning it would drop the approval columns)"
            )
        if block.editable:
            failures.append("the revision block is editable by the generation agent")
        if block.protected:
            failures.append("the revision block is protected, so publish cannot append to it")
    if schema.notes.get("house_pattern") != "ymsli_revision_block":
        failures.append("house pattern not reported on the schema")

    # All 291 paragraphs are style Normal, so the ladder is the only outline.
    if schema.notes.get("styles_used") != ["Normal"]:
        failures.append(f"styles_used={schema.notes.get('styles_used')}, want ['Normal']")
    if not schema.notes.get("heading_ladder"):
        failures.append("no heading ladder inferred from font sizes")
    titles = by_kind.get("title", [])
    if not any("Business Requirement Document" in (r.label or "") for r in titles):
        failures.append("the 24pt document title was not detected as a title")

    # The work-item table.
    tables = by_kind.get("table_data", [])
    if not any(r.notes.get("row_count") == 9 for r in tables):
        failures.append(f"the 9-row work-item table was not detected: {[r.notes.get('row_count') for r in tables]}")

    # Fields.
    fields = {r.label: r for r in by_kind.get("field", [])}
    if "Prepared by" not in fields:
        failures.append(f"'Prepared by: YMSLI' not detected as a field: {sorted(fields)}")
    elif fields["Prepared by"].notes.get("label_prefix") != "Prepared by: ":
        failures.append(
            f"label prefix not captured verbatim: {fields['Prepared by'].notes.get('label_prefix')!r}"
        )
    if "Area Path" not in fields:
        failures.append("'Area Path: ...' not detected as a field")

    # The hyperlink guard: text inside a <w:hyperlink> must never be writable.
    unsafe = [
        r
        for r in schema.regions
        if r.editable and r.notes.get("run_text_complete") is False
    ]
    if unsafe:
        failures.append(
            f"{len(unsafe)} editable region(s) hold text inside a hyperlink: "
            f"{[r.region_id for r in unsafe]}"
        )
    if schema.notes.get("hyperlink_paragraphs") != 14:
        failures.append(
            f"hyperlink_paragraphs={schema.notes.get('hyperlink_paragraphs')}, want 14"
        )

    # Fail-safe default: nothing content-bearing is left unclassified.
    covered = set()
    for region in schema.regions:
        locator = region.locator
        if locator.para_index is not None:
            covered.add(locator.para_index)
        if locator.para_span:
            covered.update(range(locator.para_span[0], locator.para_span[1] + 1))
        if region.notes.get("caption_para") is not None:
            covered.add(region.notes["caption_para"])
    document = docx.Document(io.BytesIO(data))
    uncovered = [
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() and index not in covered
    ]
    if uncovered:
        failures.append(
            f"{len(uncovered)} content paragraph(s) belong to no region: {uncovered[:12]}"
        )

    return failures


def check_synthetic() -> list[str]:
    """The two shapes the BRD has no example of.

    A `table_loop` (header row, no data) is the docx form of the MOM's Action
    Plan, and a caption with nothing under it is the form of its empty
    `Attendees` field. Neither appears in the sample, and both are writable
    kinds, so they are synthesised rather than shipped untested.
    """
    failures: list[str] = []
    document = docx.Document()
    title = document.add_paragraph()
    run = title.add_run("Workshop Readout")
    run.font.size = Pt(24)
    document.add_paragraph("Prepared by: YMSLI")
    document.add_paragraph("Action Plan:")
    table = document.add_table(rows=1, cols=3)
    for cell, name in zip(table.rows[0].cells, ("Sl.No.", "Action", "Owner")):
        cell.text = name
    document.add_paragraph("Comments:")
    buffer = io.BytesIO()
    document.save(buffer)

    schema = detect_regions_docx(buffer.getvalue())
    kinds = schema.counts_by_kind()

    loops = [r for r in schema.regions if r.kind == "table_loop"]
    if not loops:
        failures.append(f"synthetic: header-only table not detected as table_loop ({kinds})")
    else:
        loop = loops[0]
        if [c["name"] for c in loop.notes["columns"]] != ["Sl.No.", "Action", "Owner"]:
            failures.append(f"synthetic: table_loop columns wrong: {loop.notes['columns']}")
        if loop.notes["row_count"] != 0 or loop.notes["insert_at_row"] != 1:
            failures.append(f"synthetic: table_loop insertion point wrong: {loop.notes}")
        if loop.label != "Action Plan":
            failures.append(f"synthetic: table caption not attached: {loop.label!r}")
        if not loop.editable:
            failures.append("synthetic: table_loop is not editable")

    empties = {r.label: r for r in schema.regions if r.kind == "field_empty"}
    if "Comments" not in empties:
        failures.append(f"synthetic: trailing 'Comments:' not detected as field_empty ({kinds})")
    elif not empties["Comments"].editable:
        failures.append("synthetic: field_empty is not editable")

    fields = {r.label for r in schema.regions if r.kind == "field"}
    if "Prepared by" not in fields:
        failures.append(
            f"synthetic: the table below 'Prepared by: YMSLI' consumed it as a caption ({fields})"
        )

    if not any(r.kind == "title" for r in schema.regions):
        failures.append("synthetic: the 24pt line was not detected as a title")
    return failures


def main() -> int:
    path = FILES_DIR / SAMPLE
    if not path.exists():
        print(f"missing sample: {path}")
        return 1

    data = path.read_bytes()
    started = time.perf_counter()
    schema = detect_regions_docx(data)
    elapsed_ms = (time.perf_counter() - started) * 1000
    report(SAMPLE, schema, elapsed_ms)

    expected = _expected_regions(data)
    hits, total, misses = measure_recall(schema, expected)
    recall = hits / total if total else 0.0

    print()
    print("=" * 108)
    print("recall against an independently-derived ground truth")
    print("=" * 108)
    print(f"  {hits} of {total} expected regions detected with the expected kind")
    print(f"  recall = {recall:.1%}   target = {RECALL_TARGET:.0%}")
    if misses:
        print(f"  {len(misses)} miss(es):")
        for miss in misses:
            print(f"    - {miss}")

    failures = [f"BRD: {f}" for f in check_brd(schema, data)]
    failures += [f"synthetic: {f}" for f in check_synthetic()]
    if recall < RECALL_TARGET:
        failures.append(
            f"recall {recall:.1%} is below the plan's {RECALL_TARGET:.0%} target"
        )

    print()
    print("=" * 108)
    if failures:
        print(f"docx region detection: FAIL ({len(failures)})")
        for failure in failures:
            print(f"  - {failure}")
        return 1
    print(
        f"docx region detection: PASS -- recall {recall:.1%}, "
        f"{len(schema.writable())} writable regions, every content paragraph classified"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
