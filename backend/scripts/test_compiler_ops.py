"""T3 end-to-end: derive a new MOM from the real one. Read-only on the source.

    uv run python scripts/test_compiler_ops.py

This is the plan's Tier-3 case (§3.1) -- "create the next MOM for the AP
workshop from this one" -- driven entirely by `region_id`, never by hand-typed
cell refs:

  * six fields set, including one empty `Attendees` slot and one real date
  * the meeting summary block rewritten from 14 rows to 16, which pushes the
    Action Plan header and everything under it down by two
  * three Action Plan rows inserted into a table that had no data rows at all
  * both insertions in one spec, both addressed in *original* coordinates

Then it proves the result three ways: structural QA (nothing lost, merge count
changed by exactly the amount cloning row 11's `B11:G11` merge twice should
change it), functional QA (every value present at its final address), and a
second region detection pass over the output (the compiler's work is coherent
enough that the detector re-reads the Action Plan as a populated table).

It also checks the guarantees hold in the negative: writing the logo, editing
a non-editable region, and naming a region that doesn't exist all fail.
"""

import io
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402

from _artifact import write_artifact  # noqa: E402

from app.compiler import region_ops, sheet_xml as sx  # noqa: E402
from app.compiler.protect import (  # noqa: E402
    ProtectedWriteError,
    protected_ranges_from_schema,
)
from app.compiler.xlsx_surgical import (  # noqa: E402
    CompilerError,
    SetValueOp,
    compile_change_spec,
    to_excel_serial,
)
from app.ingestion.fingerprint import fingerprint_xlsx  # noqa: E402
from app.ingestion.regions import detect_regions_xlsx  # noqa: E402
from app.qa.functional import check_assertions  # noqa: E402
from app.qa.structural import diff_fingerprints  # noqa: E402

SAMPLE_PATH = Path(__file__).resolve().parent.parent.parent / "Files" / "MOM Sample.xlsx"
OUTPUT_PATH = Path(__file__).resolve().parent.parent / "scratch" / "MOM_T3_derived.xlsx"

SUMMARY_ID = "sheet1.meeting_summary"
ACTION_PLAN_ID = "sheet1.action_plan_after_current_meeting"

# Style indices that already exist in this workbook: s=7 is the bold bordered
# serial cell used by the summary block, s=31 is a plain bordered wrapped data
# cell. Without these the Action Plan's only template is its own header row,
# whose s=12 is bold -- new rows would read as repeated headings. Synthesising
# a style would mean editing styles.xml, which the compiler deliberately won't.
#
# s=22 is the one style in this workbook carrying a date number format
# (numFmtId 15, `d-mmm-yy`) and it is also bordered, unbolded and wrapped, so
# the date columns use it. Using s=31 there would put a bare 46256 on screen --
# the compiler now refuses that outright rather than letting it ship.
ACTION_PLAN_STYLES = {
    "Sl.No.": "7",
    "Action": "31",
    "Responsibility": "31",
    "Target date": "22",
    "Closure Date": "22",
    "Status": "31",
    "Remarks": "31",
}

SUMMARY_ROWS = [
    "Agenda: continuing the Accounts Payable automation thread from cycle 1, "
    "focused on invoice header capture.",
    "AP invoice header fields: vendor, invoice number, invoice date, currency and "
    "tax registration were walked through against the D365 form.",
    "Duplicate invoice detection: the combination of vendor plus invoice number is "
    "treated as the uniqueness key.",
    "Three-way match: PO, receipt and invoice tolerances were reviewed; quantity "
    "tolerance stays at zero.",
    "Non-PO invoices: routed through a separate approval workflow with cost centre "
    "coding at line level.",
    "Tax determination: GST treatment is derived from vendor registration and place "
    "of supply.",
    "Payment terms: net 45 confirmed as the default, overridable per vendor.",
    "Invoice approval hierarchy: value bands were agreed and mapped to positions "
    "rather than named users.",
    "Exception handling: held invoices surface on a workspace queue with an ageing "
    "indicator.",
    "OCR intake: header extraction confidence below 80 percent is routed for manual "
    "review.",
    "Vendor master dependency: bank detail changes require dual approval before an "
    "invoice can be paid.",
    "Accruals: month-end accrual for received-not-invoiced is generated from receipt "
    "quantities.",
    "Reporting: AP ageing and held-invoice reports were demonstrated from the "
    "standard workspace.",
    "Integration: the inbound AP header workflow is the Logic App already registered "
    "as a technical artifact.",
    "Open point: whether credit notes follow the same approval bands as invoices.",
    "Next cycle: payment run, bank reconciliation and vendor statement matching.",
]

ACTION_ROWS = [
    {
        "Sl.No.": 1,
        "Action": "Share the AP invoice header field mapping against D365 form layout",
        "Responsibility": "Ravi Kumar",
        "Target date": date(2026, 8, 22),
        "Status": "Open",
    },
    {
        "Sl.No.": 2,
        "Action": "Confirm duplicate-invoice uniqueness key with the tax team",
        "Responsibility": "Anshul Maheshwari",
        "Target date": date(2026, 8, 27),
        "Status": "Open",
    },
    {
        "Sl.No.": 3,
        "Action": "Decide credit note approval bands and circulate for sign-off",
        "Responsibility": "Shardul Pantawane",
        "Target date": date(2026, 9, 3),
        "Status": "Open",
    },
]


def build_spec(schema) -> list:
    """Every target named by region_id -- no cell ref is typed by hand."""
    return [
        region_ops.set_field(schema, "sheet1.date", date(2026, 8, 20)),
        region_ops.set_field(
            schema,
            "sheet1.purpose",
            "Workshop Session - Accounts Payable (Invoice Header Capture, Cycle 2)",
        ),
        region_ops.set_field(schema, "sheet1.prepared_by", "YMSLI"),
        region_ops.set_field(schema, "sheet1.venue", "YMSLI Faridabad - Block C"),
        region_ops.set_field(
            schema, "sheet1.attendees_ymsli", "Ravi Kumar, Shardul Pantawane, Neha Sharma"
        ),
        region_ops.set_field(
            schema, "sheet1.attendees_ymesg", "Anshul Maheshwari, Kenji Watanabe"
        ),
        region_ops.replace_block_rows(
            schema,
            SUMMARY_ID,
            [
                {"serial": index, "content": text}
                for index, text in enumerate(SUMMARY_ROWS, start=1)
            ],
        ),
        region_ops.add_table_rows(
            schema,
            ACTION_PLAN_ID,
            ACTION_ROWS,
            column_styles=ACTION_PLAN_STYLES,
            value_types={"Target date": "date", "Closure Date": "date"},
        ),
    ]


def check_guarantees(schema, original: bytes) -> list[str]:
    """The refusals matter as much as the writes."""
    failures = []
    protected = protected_ranges_from_schema(schema)
    if not any(sx.bounds(guard.ref)[1] == 1 for guard in protected):
        failures.append("no protected range covers row 1 (the logo)")

    # A raw op aimed at the logo must not survive pre-flight.
    try:
        compile_change_spec(
            original,
            [SetValueOp("xl/worksheets/sheet1.xml", "A1", "overwritten")],
            protected=protected,
        )
        failures.append("writing A1 was allowed despite branding protection")
    except ProtectedWriteError:
        pass

    # The bridge won't even build an op for a non-editable region.
    try:
        region_ops.set_field(schema, "sheet1.rich_value_a1", "overwritten")
        failures.append("set_field built an op for the branding region")
    except region_ops.RegionNotEditableError:
        pass

    # An invented region_id is refused rather than guessed at.
    try:
        region_ops.set_field(schema, "sheet1.not_a_real_region", "x")
        failures.append("unknown region_id was accepted")
    except region_ops.UnknownRegionError:
        pass

    # An unknown column inside a real region is refused too.
    try:
        region_ops.add_table_rows(schema, ACTION_PLAN_ID, [{"Nonexistent Column": "x"}])
        failures.append("unknown column name was accepted")
    except region_ops.UnknownRegionError:
        pass

    # A date written through a style with no date number format would render as
    # a bare serial; that must fail rather than ship.
    try:
        compile_change_spec(
            original,
            [
                region_ops.add_table_rows(
                    schema,
                    ACTION_PLAN_ID,
                    [{"Target date": date(2026, 8, 22)}],
                    column_styles={"Target date": "31"},  # bordered, but General
                    value_types={"Target date": "date"},
                )
            ],
        )
        failures.append("a date through a General-formatted style was allowed")
    except CompilerError as exc:
        if "22" not in str(exc):
            failures.append(f"date-style error should name the usable style: {exc}")

    return failures


def check_merge_shifting() -> list[str]:
    """`shift_merges` directly: this sample has no merges below row 24, so the
    T3 spec never exercises the move-and-straddle branches."""
    failures = []
    moved = sx.shift_merges(["A1:G2", "B24:G24", "B30:C31"], at_row=25, delta=2)
    if moved != ["A1:G2", "B24:G24", "B32:C33"]:
        failures.append(f"shift_merges moved the wrong ranges: {moved}")
    straddled = sx.shift_merges(["B10:B20"], at_row=15, delta=3)
    if straddled != ["B10:B23"]:
        failures.append(f"a merge spanning the insertion point should grow: {straddled}")
    return failures


def check_output_geometry(schema, data: bytes) -> list[str]:
    """Where things ended up, computed independently of the compiler."""
    failures = []
    output_schema = detect_regions_xlsx(data)

    summary = next(
        (r for r in output_schema.regions if r.kind == "repeating_block"), None
    )
    if summary is None:
        failures.append("summary block no longer detected in the output")
    else:
        if summary.notes["row_count"] != len(SUMMARY_ROWS):
            failures.append(
                f"summary row_count={summary.notes['row_count']}, want {len(SUMMARY_ROWS)}"
            )
        if summary.locator.ref != "A11:G26":
            failures.append(f"summary block at {summary.locator.ref}, want A11:G26")
        if summary.notes["serial_gaps"]:
            failures.append(f"rewritten summary should renumber cleanly: {summary.notes['serial_gaps']}")

    # The Action Plan had zero data rows; now it has three, so the detector
    # should read it as a populated table rather than an empty loop.
    action = next(
        (r for r in output_schema.regions if r.label == "Action Plan after Current Meeting"),
        None,
    )
    if action is None:
        failures.append("Action Plan region no longer detected in the output")
    else:
        if action.kind != "table_data":
            failures.append(f"Action Plan detected as {action.kind}, want table_data")
        if action.notes["row_count"] != len(ACTION_ROWS):
            failures.append(
                f"Action Plan row_count={action.notes['row_count']}, want {len(ACTION_ROWS)}"
            )
        if action.notes["header_ref"] != "A29:G29":
            failures.append(
                f"Action Plan header at {action.notes['header_ref']}, want A29:G29 "
                "(27 + 2 inserted summary rows)"
            )

    if any(r.kind == "branding" and r.locator.ref == "A1:G2" for r in output_schema.regions):
        pass
    else:
        failures.append("branding region A1:G2 missing from the output schema")

    return failures


def check_table_data_replace() -> list[str]:
    """A second workbook, and the one case the MOM can't cover.

    A `repeating_block`'s locator holds only data rows; a `table_data`'s holds
    its header row too. Replacing BFL's Fit/Gap rows must leave row 2's column
    headings alone, and it exercises the compiler against a four-sheet file
    whose data sheet claims a million rows.
    """
    failures = []
    path = SAMPLE_PATH.with_name("BFL Sample.xlsx")
    if not path.exists():
        return [f"missing sample: {path}"]

    original = path.read_bytes()
    schema = detect_regions_xlsx(original)
    region = next(r for r in schema.regions if r.kind == "table_data")

    op = region_ops.replace_block_rows(
        schema,
        region.region_id,
        [
            {"ID": 900, "Business process": "65.90 Accounts Payable Automation"},
            {"ID": 901, "Business process": "65.90.005 Capture AP invoice header"},
        ],
    )
    if op.ref.startswith("A2"):
        failures.append(f"replace range {op.ref} includes the header row")

    result = compile_change_spec(
        original, [op], protected=protected_ranges_from_schema(schema)
    )
    try:
        diff_fingerprints(
            fingerprint_xlsx(original), fingerprint_xlsx(result.data), expected_merge_delta=0
        )
        check_assertions(result.data, result.assertions)
    except Exception as exc:  # noqa: BLE001
        failures.append(f"BFL replace: {exc}")

    sheet = openpyxl.load_workbook(io.BytesIO(result.data), data_only=True)[
        region.locator.sheet
    ]
    headers = [sheet.cell(row=2, column=col).value for col in (1, 2, 3)]
    if headers != ["ID", "Level", "Business process"]:
        failures.append(f"BFL header row was overwritten: {headers}")
    if sheet["A3"].value != 900 or sheet["C4"].value != "65.90.005 Capture AP invoice header":
        failures.append(
            f"BFL rows not written: A3={sheet['A3'].value!r} C4={sheet['C4'].value!r}"
        )
    if sheet["C5"].value is not None:
        failures.append(f"BFL leftover row 5 not blanked: {sheet['C5'].value!r}")
    return failures


def main() -> int:
    if not SAMPLE_PATH.exists():
        print(f"sample not found: {SAMPLE_PATH}")
        return 1

    original = SAMPLE_PATH.read_bytes()
    schema = detect_regions_xlsx(original)
    before = fingerprint_xlsx(original)
    print(
        f"source: {len(before.package_parts)} parts, {before.merged_ranges} merges, "
        f"{len(schema.regions)} regions ({len(schema.writable())} writable)"
    )

    failures = check_guarantees(schema, original) + check_merge_shifting()

    spec = build_spec(schema)
    result = compile_change_spec(
        original, spec, protected=protected_ranges_from_schema(schema)
    )
    after = fingerprint_xlsx(result.data)
    print(
        f"compiled: {len(spec)} ops -> {result.operations_applied} primitives, "
        f"{result.rows_inserted} rows inserted, merge delta {result.merge_delta:+d}, "
        f"{len(result.assertions)} assertions"
    )

    # Two rows cloned from row 11, whose B11:G11 merge is cloned with them.
    expected_merge_delta = 2
    if result.merge_delta != expected_merge_delta:
        failures.append(
            f"merge delta {result.merge_delta}, want {expected_merge_delta} "
            "(one per cloned summary row)"
        )

    try:
        structural = diff_fingerprints(before, after, expected_merge_delta=expected_merge_delta)
        print(
            f"structural QA: PASS -- parts lost {structural.lost_parts}, "
            f"merges {structural.merged_ranges_before}->{structural.merged_ranges_after}, "
            f"label {after.sensitivity_label}, richData {after.rich_value_parts}, "
            f"media {after.media_files}"
        )
    except Exception as exc:  # noqa: BLE001 - report, don't mask
        failures.append(f"structural QA: {exc}")

    try:
        functional = check_assertions(result.data, result.assertions)
        print(f"functional QA: PASS -- {len(functional.checks)} values verified in place")
    except Exception as exc:  # noqa: BLE001
        failures.append(f"functional QA: {exc}")

    failures += check_output_geometry(schema, result.data)
    failures += check_table_data_replace()

    print()
    print("landmarks in the derived document:")
    for label, ref, expected in (
        ("date field", "B5", to_excel_serial(date(2026, 8, 20))),
        ("attendees (was empty)", "B7", None),
        ("summary row 16 (inserted)", "B26", None),
        ("Action Plan header (was row 27)", "A29", None),
        ("first action row (inserted)", "B30", None),
    ):
        print(f"  {label:<34} {ref}")

    print(f"\nwrote {write_artifact(OUTPUT_PATH, result.data)}")

    if failures:
        print(f"\ncompiler ops: FAIL ({len(failures)})")
        for failure in failures:
            print(f"  - {failure}")
        return 1
    print("\ncompiler ops: PASS -- open the file in real Excel to confirm the logo and borders")
    return 0


if __name__ == "__main__":
    sys.exit(main())
