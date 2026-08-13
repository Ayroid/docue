"""Build compiler operations from detected regions.

This is the injection barrier from plan §3.2: *the model may only emit
`region_id`s present in the schema; unknown targets are dropped, not guessed.*
Everything upstream (a question plan, an LLM, the API) names a `region_id` and
supplies values. Only this module turns a name into a cell reference, so no
caller is ever in a position to invent a ref, a sheet part or a style index.

It also refuses to build an operation for a region the detector marked
non-editable, which means branding, formulas and unclassified cells are
unreachable from here by construction -- `protect.py` is the second line of
defence, not the only one.
"""

from typing import Any

from openpyxl.utils import get_column_letter

from app.compiler import sheet_xml as sx
from app.compiler.xlsx_surgical import (
    ClearRegionOp,
    CloneTableRowsOp,
    ReplaceRegionRowsOp,
    SetValueOp,
    ValueType,
    infer_value_type,
)


class UnknownRegionError(KeyError):
    """No such region in this document's schema."""


class RegionNotEditableError(RuntimeError):
    """The region exists but the detector marked it non-writable."""


def find_region(schema: Any, region_id: str) -> Any:
    for region in schema.regions:
        if region.region_id == region_id:
            return region
    raise UnknownRegionError(region_id)


def _editable(region: Any) -> Any:
    if not region.editable:
        raise RegionNotEditableError(
            f"{region.region_id} is {region.kind} (editable=False, "
            f"protected={region.protected})"
        )
    return region


def _sheet_file(region: Any) -> str:
    sheet_file = region.locator.sheet_file
    if sheet_file is None:
        raise RegionNotEditableError(f"{region.region_id} has no worksheet part")
    return sheet_file


def anchor_ref(ref: str) -> str:
    """Top-left cell of a range: the only cell of a merge that holds a value."""
    min_col, min_row, _, _ = sx.bounds(ref)
    return f"{get_column_letter(min_col)}{min_row}"


def column_map(region: Any) -> dict[str, str]:
    """Caller-facing column names -> column letters, taken from the schema.

    Tables are named by their header text. A repeating block has no header row
    to name, so it exposes `serial` and `content` (plus whatever its serial
    column was captioned, e.g. MOM's `Sl.No.`).
    """
    notes = region.notes
    if region.kind == "repeating_block":
        mapping = {
            "serial": notes["serial_column"],
            "content": notes["content_column"],
        }
        if notes.get("column_header"):
            mapping[notes["column_header"]] = notes["serial_column"]
        return mapping
    return {column["name"]: column["column"] for column in notes.get("columns", [])}


def _translate(
    rows: list[dict[str, Any]], mapping: dict[str, str], region_id: str, strict: bool
) -> list[dict[str, Any]]:
    translated: list[dict[str, Any]] = []
    for row in rows:
        converted: dict[str, Any] = {}
        for name, value in row.items():
            column = mapping.get(name) or (name if name in mapping.values() else None)
            if column is None:
                if strict:
                    raise UnknownRegionError(
                        f"{region_id} has no column {name!r}; known: {sorted(mapping)}"
                    )
                continue  # drop, never guess
            converted[column] = value
        translated.append(converted)
    return translated


def _value_types(
    value_types: dict[str, ValueType] | None, mapping: dict[str, str]
) -> dict[str, ValueType]:
    if not value_types:
        return {}
    return {mapping.get(name, name): kind for name, kind in value_types.items()}


def set_field(
    schema: Any, region_id: str, value: Any, *, value_type: ValueType | None = None
) -> SetValueOp:
    """Write a `field` / `field_empty` region's value slot.

    The region's own `value_type` wins when the detector could read one, since
    it reflects what the template already holds. An empty field has nothing to
    infer from, so the Python value decides -- a `date` there is a date.
    """
    region = _editable(find_region(schema, region_id))
    if value_type is None:
        value_type = (
            region.value_type
            if region.value_type in ("string", "number", "date")
            else infer_value_type(value)
        )
    return SetValueOp(
        sheet_file=_sheet_file(region),
        ref=anchor_ref(region.locator.ref),
        value=value,
        value_type=value_type,
    )


def clear_region(schema: Any, region_id: str) -> ClearRegionOp:
    region = _editable(find_region(schema, region_id))
    return ClearRegionOp(sheet_file=_sheet_file(region), ref=region.locator.ref)


def add_table_rows(
    schema: Any,
    region_id: str,
    rows: list[dict[str, Any]],
    *,
    column_styles: dict[str, str] | None = None,
    value_types: dict[str, ValueType] | None = None,
    strict: bool = True,
) -> CloneTableRowsOp:
    """Append rows to a `table_loop` / `table_data` region.

    A `table_loop` has no data row to copy styling from, so the header row is
    the only template available; pass `column_styles` (indices that already
    exist in the workbook) to override where the header's own styling would
    look wrong, e.g. bold cells for what should be plain data.
    """
    region = _editable(find_region(schema, region_id))
    notes = region.notes
    mapping = column_map(region)
    header_row = sx.bounds(notes["header_ref"])[1]
    _, _, _, max_row = sx.bounds(region.locator.ref)

    return CloneTableRowsOp(
        sheet_file=_sheet_file(region),
        at_row=notes.get("insert_at_row", max_row + 1),
        rows=_translate(rows, mapping, region_id, strict),
        row_template_row=notes.get("first_data_row") if notes.get("row_count") else header_row,
        column_styles=_translate([column_styles or {}], mapping, region_id, strict)[0],
        value_types=_value_types(value_types, mapping),
    )


def replace_block_rows(
    schema: Any,
    region_id: str,
    rows: list[dict[str, Any]],
    *,
    column_styles: dict[str, str] | None = None,
    value_types: dict[str, ValueType] | None = None,
    strict: bool = True,
) -> ReplaceRegionRowsOp:
    """Rewrite a block's rows in full, growing or blanking it to fit.

    A `repeating_block`'s locator covers only data rows, but a table's covers
    its header row too -- so the replaceable range starts at `first_data_row`
    where the schema records one. Writing from the locator's first row instead
    would overwrite the column headings with data.
    """
    region = _editable(find_region(schema, region_id))
    mapping = column_map(region)
    min_col, min_row, max_col, max_row = sx.bounds(region.locator.ref)
    first_data_row = region.notes.get("first_data_row", min_row)

    if first_data_row > max_row:
        raise RegionNotEditableError(
            f"{region_id} has no data rows to replace (it is {region.kind}); "
            "use add_table_rows to populate it"
        )

    return ReplaceRegionRowsOp(
        sheet_file=_sheet_file(region),
        ref=(
            f"{get_column_letter(min_col)}{first_data_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        ),
        rows=_translate(rows, mapping, region_id, strict),
        row_template_row=first_data_row,
        column_styles=_translate([column_styles or {}], mapping, region_id, strict)[0],
        value_types=_value_types(value_types, mapping),
    )
