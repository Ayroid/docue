"""pptx compiler: shape-level edits through python-pptx.

This is the opposite strategy to the xlsx compiler, and the reason is a
measurement rather than a preference. openpyxl's round-trip drops nine parts
from the MOM including the logo, so xlsx is edited surgically at the zip level.
python-pptx round-trips the corporate deck with all 52 parts intact and the
Purview label byte-identical, and the generated template with all 61 parts
byte-identical -- chart XML and the chart's embedded workbook included. So here
the object model is safe to use, and using it buys shape ids, placeholder types
and table access for free.

What still needs care is *formatting*. python-pptx's convenient setters throw it
away:

    shape.text_frame.text = "new"      # every run property on the shape, gone
    cell.text = "new"                  # same for a table cell

Both replace the whole text body with one unformatted run. So every write here
goes through `_replace_paragraph_text`, which keeps the first run's `<a:rPr>`
and drops the rest, and every new paragraph or table row is a deepcopy of an
existing one -- inheriting bullet level, indentation, font and row height
instead of guessing at them. Same principle as the xlsx compiler reusing style
indices rather than writing `styles.xml`.
"""

import copy
import io
from dataclasses import dataclass, field
from typing import Any, Union

from pptx import Presentation
from pptx.oxml.ns import qn

_RUN_TAGS = ("a:r", "a:br", "a:fld")


class PptxCompilerError(RuntimeError):
    """The spec cannot be compiled without risking damage to the deck."""


class ProtectedShapeError(PptxCompilerError):
    """An operation targets a shape that must not be written."""


@dataclass
class SetTextOp:
    """Replace a shape's text, keeping its first run's formatting."""

    slide_part: str  # "ppt/slides/slide1.xml"
    shape_id: int
    text: str


@dataclass
class SetBulletsOp:
    """Replace a bullet list, cloning one paragraph's formatting per item."""

    slide_part: str
    shape_id: int
    items: list[str]
    template_paragraph: int = 0


@dataclass
class SetTableRowsOp:
    """Rewrite a table's data rows: fill what exists, clone for the rest.

    `rows` are keyed by column index. A template deck ships its action table
    with a header and one empty row precisely so there is a row to clone; that
    empty row is filled first rather than left stranded above the data.
    """

    slide_part: str
    shape_id: int
    rows: list[dict[int, Any]]
    first_data_row: int = 1
    template_row: int | None = None  # defaults to the last existing row


@dataclass
class SetNotesOp:
    slide_part: str
    text: str


@dataclass
class ProtectShapeOp:
    """A declaration, not an edit: no other op may write this shape."""

    slide_part: str
    shape_id: int
    reason: str = ""


Operation = Union[SetTextOp, SetBulletsOp, SetTableRowsOp, SetNotesOp, ProtectShapeOp]


@dataclass
class PptxAssertion:
    slide_part: str
    shape_id: int | None
    expected: str
    where: str  # "text" | "bullet:N" | "cell:R,C" | "notes"


@dataclass
class PptxCompileResult:
    data: bytes
    operations_applied: int
    rows_added: int
    paragraphs_written: int
    assertions: list[PptxAssertion]


@dataclass(frozen=True)
class ProtectedShape:
    slide_part: str
    shape_id: int
    reason: str = ""
    region_id: str | None = None


# ── protection ────────────────────────────────────────────────────────────


def protected_shapes_from_schema(schema: Any) -> list[ProtectedShape]:
    """Lift every protected region out of a `PptxDocumentSchema`."""
    shapes = []
    for region in getattr(schema, "regions", []):
        if not region.protected or region.locator.shape_id is None:
            continue
        shapes.append(
            ProtectedShape(
                slide_part=region.locator.slide_part,
                shape_id=region.locator.shape_id,
                reason=region.kind,
                region_id=region.region_id,
            )
        )
    return shapes


def preflight(operations: list[Operation], protected: list[ProtectedShape]) -> None:
    """Raise before anything is touched if an op would write a protected shape."""
    guarded = {(shape.slide_part, shape.shape_id): shape for shape in protected}
    for op in operations:
        if isinstance(op, ProtectShapeOp):
            guarded[(op.slide_part, op.shape_id)] = ProtectedShape(
                op.slide_part, op.shape_id, op.reason or "declared"
            )
    if not guarded:
        return

    for op in operations:
        if isinstance(op, (ProtectShapeOp, SetNotesOp)):
            continue
        shape = guarded.get((op.slide_part, op.shape_id))
        if shape is not None:
            raise ProtectedShapeError(
                f"{type(op).__name__} targets shape {op.shape_id} on "
                f"{op.slide_part}, which is protected "
                f"({shape.region_id or shape.reason})"
            )


# ── text primitives ───────────────────────────────────────────────────────


def _replace_paragraph_text(paragraph: Any, text: str) -> None:
    """Set a paragraph's text, keeping the first run's character formatting.

    Runs beyond the first are removed along with any `<a:br>` or `<a:fld>`,
    which would otherwise leave fragments of the old text behind.
    """
    element = paragraph._p
    runs = [
        child
        for child in element
        if child.tag in {qn(tag) for tag in _RUN_TAGS}
    ]

    if not text:
        for run in runs:
            element.remove(run)
        return

    if not runs:
        paragraph.add_run().text = text
        return

    keep = runs[0]
    for run in runs[1:]:
        element.remove(run)
    if keep.tag != qn("a:r"):
        # A lone <a:br> or <a:fld> is not a run we can write into.
        element.remove(keep)
        paragraph.add_run().text = text
        return
    text_element = keep.find(qn("a:t"))
    if text_element is None:
        text_element = keep.makeelement(qn("a:t"), {})
        keep.append(text_element)
    text_element.text = text


def _set_shape_text(shape: Any, text: str) -> None:
    frame = shape.text_frame
    _replace_paragraph_text(frame.paragraphs[0], text)
    body = frame._txBody
    for paragraph in body.findall(qn("a:p"))[1:]:
        body.remove(paragraph)


def _set_bullets(shape: Any, items: list[str], template_index: int) -> int:
    frame = shape.text_frame
    paragraphs = frame.paragraphs
    if template_index >= len(paragraphs):
        raise PptxCompilerError(
            f"paragraph {template_index} does not exist on shape {shape.shape_id}"
        )
    template = copy.deepcopy(paragraphs[template_index]._p)

    body = frame._txBody
    for paragraph in body.findall(qn("a:p")):
        body.remove(paragraph)
    for item in items:
        body.append(copy.deepcopy(template))
        _replace_paragraph_text(frame.paragraphs[-1], item)
    return len(items)


def _set_cell_text(cell: Any, text: str) -> None:
    frame = cell.text_frame
    _replace_paragraph_text(frame.paragraphs[0], text)
    body = frame._txBody
    for paragraph in body.findall(qn("a:p"))[1:]:
        body.remove(paragraph)


# ── table primitives ──────────────────────────────────────────────────────


def _set_table_rows(shape: Any, op: SetTableRowsOp) -> int:
    if not getattr(shape, "has_table", False):
        raise PptxCompilerError(f"shape {op.shape_id} on {op.slide_part} is not a table")
    table = shape.table
    element = table._tbl
    existing = len(table.rows)
    template_index = op.template_row if op.template_row is not None else existing - 1
    if not 0 <= template_index < existing:
        raise PptxCompilerError(
            f"row template {template_index} is outside the table (rows={existing})"
        )
    if template_index < op.first_data_row and existing > op.first_data_row:
        raise PptxCompilerError(
            f"row template {template_index} is a header row; cloning it would "
            "repeat the headings as data"
        )

    template_row = copy.deepcopy(element.findall(qn("a:tr"))[template_index])
    capacity = existing - op.first_data_row
    added = max(0, len(op.rows) - capacity)
    for _ in range(added):
        element.append(copy.deepcopy(template_row))
        shape.height += table.rows[len(table.rows) - 1].height

    column_count = len(table.columns)
    for index in range(max(len(op.rows), capacity)):
        row_index = op.first_data_row + index
        if row_index >= len(table.rows):
            break
        values = op.rows[index] if index < len(op.rows) else {}
        for column in range(column_count):
            # Cells with no value are cleared rather than left holding whatever
            # the cloned template row carried.
            _set_cell_text(table.cell(row_index, column), str(values.get(column, "")))
    return added


# ── compilation ───────────────────────────────────────────────────────────


def _slides_by_part(presentation: Any) -> dict[str, Any]:
    return {
        str(slide.part.partname).lstrip("/"): slide for slide in presentation.slides
    }


def _shape_by_id(slide: Any, shape_id: int, slide_part: str) -> Any:
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise PptxCompilerError(f"no shape with id {shape_id} on {slide_part}")


def compile_pptx_spec(
    original_bytes: bytes,
    operations: list[Operation],
    *,
    protected: list[ProtectedShape] | None = None,
) -> PptxCompileResult:
    """Apply a Change Spec to a deck and report what it wrote."""
    preflight(operations, protected or [])

    presentation = Presentation(io.BytesIO(original_bytes))
    slides = _slides_by_part(presentation)

    assertions: list[PptxAssertion] = []
    rows_added = 0
    paragraphs_written = 0
    applied = 0

    for op in operations:
        if isinstance(op, ProtectShapeOp):
            continue
        slide = slides.get(op.slide_part)
        if slide is None:
            raise PptxCompilerError(f"no slide part {op.slide_part} in this deck")

        if isinstance(op, SetNotesOp):
            frame = slide.notes_slide.notes_text_frame
            _replace_paragraph_text(frame.paragraphs[0], op.text)
            body = frame._txBody
            for paragraph in body.findall(qn("a:p"))[1:]:
                body.remove(paragraph)
            assertions.append(PptxAssertion(op.slide_part, None, op.text, "notes"))
            applied += 1
            continue

        shape = _shape_by_id(slide, op.shape_id, op.slide_part)

        if isinstance(op, SetTextOp):
            _set_shape_text(shape, op.text)
            assertions.append(
                PptxAssertion(op.slide_part, op.shape_id, op.text, "text")
            )
        elif isinstance(op, SetBulletsOp):
            paragraphs_written += _set_bullets(shape, op.items, op.template_paragraph)
            for index, item in enumerate(op.items):
                assertions.append(
                    PptxAssertion(op.slide_part, op.shape_id, item, f"bullet:{index}")
                )
        elif isinstance(op, SetTableRowsOp):
            rows_added += _set_table_rows(shape, op)
            for index, values in enumerate(op.rows):
                row_index = op.first_data_row + index
                for column, value in values.items():
                    assertions.append(
                        PptxAssertion(
                            op.slide_part,
                            op.shape_id,
                            str(value),
                            f"cell:{row_index},{column}",
                        )
                    )
        else:  # pragma: no cover - the Union is closed
            raise PptxCompilerError(f"unsupported operation {type(op).__name__}")
        applied += 1

    buffer = io.BytesIO()
    presentation.save(buffer)
    return PptxCompileResult(
        data=buffer.getvalue(),
        operations_applied=applied,
        rows_added=rows_added,
        paragraphs_written=paragraphs_written,
        assertions=assertions,
    )


# ── from the region schema ────────────────────────────────────────────────
#
# The same injection barrier as the xlsx side: callers name a region_id, and
# only this code turns one into a slide part and a shape id. Unknown ids and
# non-editable regions are refused, never guessed at.


class UnknownRegionError(KeyError):
    """No such region in this deck's schema."""


class RegionNotEditableError(PptxCompilerError):
    """The region exists but the detector marked it non-writable."""


def find_region(schema: Any, region_id: str) -> Any:
    for region in schema.regions:
        if region.region_id == region_id:
            return region
    raise UnknownRegionError(region_id)


def _editable(region: Any) -> Any:
    if not region.editable:
        raise RegionNotEditableError(
            f"{region.region_id} is {region.kind} (editable=False, "
            f"protected={region.protected})"
        )
    return region


def set_text(schema: Any, region_id: str, text: str) -> SetTextOp | SetNotesOp:
    """Write a `title` / `field` / `field_empty` region, or speaker notes."""
    region = _editable(find_region(schema, region_id))
    if region.locator.notes:
        return SetNotesOp(slide_part=region.locator.slide_part, text=text)
    return SetTextOp(
        slide_part=region.locator.slide_part,
        shape_id=region.locator.shape_id,
        text=text,
    )


def set_bullets(schema: Any, region_id: str, items: list[str]) -> SetBulletsOp:
    """Rewrite a `repeating_block` -- a bullet list of any length."""
    region = _editable(find_region(schema, region_id))
    if region.kind != "repeating_block":
        raise RegionNotEditableError(
            f"{region_id} is {region.kind}, not a bullet list; use set_text"
        )
    return SetBulletsOp(
        slide_part=region.locator.slide_part,
        shape_id=region.locator.shape_id,
        items=items,
    )


def set_table_rows(
    schema: Any,
    region_id: str,
    rows: list[dict[str, Any]],
    *,
    strict: bool = True,
) -> SetTableRowsOp:
    """Rewrite a table region's data rows, addressed by column heading."""
    region = _editable(find_region(schema, region_id))
    if region.kind not in ("table_loop", "table_data"):
        raise RegionNotEditableError(f"{region_id} is {region.kind}, not a table")

    by_name = {column["name"]: column["index"] for column in region.notes["columns"]}
    translated: list[dict[int, Any]] = []
    for row in rows:
        converted: dict[int, Any] = {}
        for name, value in row.items():
            index = by_name.get(name)
            if index is None:
                if strict:
                    raise UnknownRegionError(
                        f"{region_id} has no column {name!r}; known: {sorted(by_name)}"
                    )
                continue  # drop, never guess
            converted[index] = value
        translated.append(converted)

    return SetTableRowsOp(
        slide_part=region.locator.slide_part,
        shape_id=region.locator.shape_id,
        rows=translated,
        first_data_row=region.notes["first_data_row"],
        template_row=region.notes["row_template_row"],
    )
