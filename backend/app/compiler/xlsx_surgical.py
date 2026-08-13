"""Surgical zip-level xlsx compiler.

openpyxl round-trips (load_workbook -> save) silently drop package parts it has
no model for -- in-cell rich-value images, docMetadata sensitivity labels,
richData. This module never calls openpyxl.save(). It rewrites only the
worksheet parts named in the Change Spec and copies every other zip part
byte-for-byte from the original archive, so nothing the library doesn't
understand can be lost.

Operations mirror the Change Specification in plan §3.2: `set_value`,
`clear_region`, `clone_table_rows`, `replace_region_rows` and the declarative
`protect`. Three properties hold across a whole spec:

- **Every ref is in original-document coordinates.** A spec that writes `E5`
  and also inserts rows at 25 doesn't have to know that `E5` is unaffected or
  that a later Action Plan row moved. Value writes are applied before
  insertions, and insertions are applied bottom-up, so the coordinates a
  caller read off the region schema are the coordinates that work.
- **The result reports where things landed.** `CompileResult.assertions` gives
  final-coordinate refs and expected values, which is the functional QA pass
  in plan §11 without any extra bookkeeping.
- **Styles are reused, never invented.** New rows borrow style indices that
  already exist in the workbook. Synthesising a style would mean editing
  `styles.xml`, and the fewer parts this compiler touches the stronger the
  preservation guarantee.
"""

import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Literal, Union

from openpyxl.utils import column_index_from_string, get_column_letter

from app.compiler import sheet_xml as sx
from app.compiler.styles import StyleFormats, read_style_formats

ValueType = Literal["string", "number", "date"]

# Serial 1 is 1900-01-01, but Excel also believes 1900-02-29 existed, so the
# arithmetic epoch is two days before 1900-01-01 rather than one.
_EPOCH_1900 = date(1899, 12, 30)
_EPOCH_1904 = date(1904, 1, 1)


class CompilerError(RuntimeError):
    """The spec cannot be compiled without risking corruption."""


@dataclass
class SetValueOp:
    """Write one cell, keeping its existing style index (and so its format)."""

    sheet_file: str  # e.g. "xl/worksheets/sheet1.xml"
    ref: str  # e.g. "E5"
    value: Any
    value_type: ValueType = "string"
    preserve_number_format: bool = True  # keep the cell's existing `s` index
    style: str | None = None  # only used when the cell has no style of its own


@dataclass
class ClearRegionOp:
    """Blank every cell in a range, keeping styles and merges intact."""

    sheet_file: str
    ref: str  # "A15:G24"


@dataclass
class CloneTableRowsOp:
    """Insert rows, cloning an existing row's styling, heights and merges.

    `rows` are keyed by column letter. The row at `row_template_row` supplies
    per-column style indices; `column_styles` overrides those per column with
    another index that already exists in the workbook.
    """

    sheet_file: str
    at_row: int
    rows: list[dict[str, Any]]
    row_template_row: int
    column_styles: dict[str, str] = field(default_factory=dict)
    value_types: dict[str, ValueType] = field(default_factory=dict)


@dataclass
class ReplaceRegionRowsOp:
    """Rewrite a block's rows: fill what exists, then grow or blank the rest."""

    sheet_file: str
    ref: str  # the block as detected, e.g. "A11:G24"
    rows: list[dict[str, Any]]
    row_template_row: int | None = None  # defaults to the block's first row
    column_styles: dict[str, str] = field(default_factory=dict)
    value_types: dict[str, ValueType] = field(default_factory=dict)


@dataclass
class ProtectOp:
    """A declaration, not an edit: no other op may write inside this range."""

    sheet_file: str
    ref: str
    reason: str = ""


Operation = Union[
    SetValueOp, ClearRegionOp, CloneTableRowsOp, ReplaceRegionRowsOp, ProtectOp
]


@dataclass
class Assertion:
    """An expected value at a *final* ref -- feeds functional QA."""

    sheet_file: str
    ref: str
    expected: Any
    value_type: ValueType


@dataclass
class CompileResult:
    data: bytes
    rows_inserted: int
    merge_delta: int
    operations_applied: int
    assertions: list[Assertion]


# ── value encoding ────────────────────────────────────────────────────────


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_excel_serial(value: date | datetime | str | int | float, *, date1904: bool = False) -> float:
    """Excel stores dates as numbers; the number format makes them look like
    dates. Write the serial and keep the format, or `d-mmm-yy` renders 46250.
    """
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if isinstance(value, str):
        value = datetime.fromisoformat(value)
    epoch = _EPOCH_1904 if date1904 else _EPOCH_1900
    if isinstance(value, datetime):
        whole = (value.date() - epoch).days
        fraction = (value.hour * 3600 + value.minute * 60 + value.second) / 86400
        return whole + fraction
    return float((value - epoch).days)


def infer_value_type(value: Any) -> ValueType:
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, bool):
        return "string"
    if isinstance(value, (int, float)):
        return "number"
    return "string"


def build_cell(
    ref: str,
    value: Any,
    value_type: ValueType,
    style: str | None,
    *,
    date1904: bool = False,
) -> str:
    style_attr = f' s="{style}"' if style else ""
    if value is None or value == "":
        return f'<c r="{ref}"{style_attr}/>'
    if value_type == "string":
        text = _xml_escape(str(value))
        # inlineStr keeps sharedStrings.xml out of it entirely -- one less part
        # to keep consistent, and no bookkeeping to get wrong.
        return f'<c r="{ref}"{style_attr} t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>'
    if value_type == "date":
        return f'<c r="{ref}"{style_attr}><v>{to_excel_serial(value, date1904=date1904)}</v></c>'
    return f'<c r="{ref}"{style_attr}><v>{value}</v></c>'


# ── per-sheet edit session ────────────────────────────────────────────────


class _SheetEditor:
    """Accumulates edits to one worksheet part, then renders it once."""

    def __init__(self, xml: str, *, date1904: bool, formats: StyleFormats) -> None:
        self.rows: dict[int, sx.Row] = {row.number: row for row in sx.read_rows(xml)}
        self.merges: list[str] = sx.read_merges(xml)
        self.original_merges = list(self.merges)
        self.xml = xml
        self.date1904 = date1904
        self.formats = formats
        self.rows_inserted = 0
        self.merges_before = len(self.merges)
        self.max_col_touched = 0

    def _check_date_style(self, ref: str, value_type: ValueType, style: str | None) -> None:
        """A date written through a General-formatted style renders as 46256."""
        if value_type != "date" or self.formats.renders_as_date(style):
            return
        usable = self.formats.date_styles()
        raise CompilerError(
            f"{ref} would write a date through style s={style or '0'} "
            f"(numFmtId {self.formats.numfmt_of(style)}), which renders the serial "
            f"as a number. Styles in this workbook that render as dates: "
            f"{usable or 'none -- styles.xml has no date format to reuse'}"
        )

    # -- reads ---------------------------------------------------------

    def cell(self, ref: str) -> str | None:
        row = self.rows.get(sx.ref_row(ref))
        if row is None:
            return None
        for cell in sx.read_cells(row.text):
            if sx.cell_ref(cell) == ref:
                return cell
        return None

    def style_at(self, ref: str) -> str | None:
        cell = self.cell(ref)
        return sx.cell_style(cell) if cell else None

    def row_styles(self, row_number: int) -> dict[str, str]:
        row = self.rows.get(row_number)
        if row is None:
            return {}
        styles = {}
        for cell in sx.read_cells(row.text):
            style = sx.cell_style(cell)
            if style is not None:
                styles[sx.ref_column(sx.cell_ref(cell))] = style
        return styles

    # -- writes --------------------------------------------------------

    def set_value(self, op: SetValueOp) -> None:
        row_number = sx.ref_row(op.ref)
        row = self.rows.get(row_number)
        style = (self.style_at(op.ref) if op.preserve_number_format else None) or op.style
        self._check_date_style(op.ref, op.value_type, style)
        cell_xml = build_cell(op.ref, op.value, op.value_type, style, date1904=self.date1904)
        column = sx.ref_column(op.ref)
        if row is None:
            # A cell absent from the XML is normal in a sparse sheet; create the
            # row so `set_value` isn't limited to cells Excel happened to write.
            row = sx.Row(number=row_number, text=f'<row r="{row_number}"></row>')
        edited = sx.Row(number=row_number, text=sx.set_cell(row.text, cell_xml))
        # Only touch `spans` if the write actually falls outside it.
        if not sx.covered_by_spans(row.text, column):
            edited = sx.set_row_spans(edited)
        self.rows[row_number] = edited
        self.max_col_touched = max(self.max_col_touched, column_index_from_string(column))

    def clear(self, ref: str) -> None:
        min_col, min_row, max_col, max_row = sx.bounds(ref)
        for row_number in range(min_row, max_row + 1):
            row = self.rows.get(row_number)
            if row is None:
                continue
            cells = []
            for cell in sx.read_cells(row.text):
                column = column_index_from_string(sx.ref_column(sx.cell_ref(cell)))
                if not min_col <= column <= max_col:
                    cells.append(cell)
                    continue
                if sx.cell_has_formula(cell):
                    raise CompilerError(
                        f"{sx.cell_ref(cell)} holds a formula; clearing it would "
                        f"delete a derived value. Exclude it from {ref}."
                    )
                cells.append(sx.blank_cell(cell))
            self.rows[row_number] = sx.Row(
                number=row_number, text=sx.assemble_row(sx.row_open_tag(row.text), cells)
            )

    def insert_rows(self, op: CloneTableRowsOp) -> list[tuple[int, dict[str, Any]]]:
        """Insert `op.rows` at `op.at_row`. Returns (row number, payload) pairs."""
        count = len(op.rows)
        if count == 0:
            return []

        template = self.rows.get(op.row_template_row)
        if template is None:
            raise CompilerError(
                f"row template {op.row_template_row} does not exist in {op.sheet_file}"
            )

        # Formula refs would have to be rewritten to stay correct, and that
        # needs a formula parser. Refuse rather than quietly break them.
        for row in self.rows.values():
            if row.number >= op.at_row and sx.row_has_formula(row):
                raise CompilerError(
                    f"cannot insert at row {op.at_row}: row {row.number} of "
                    f"{op.sheet_file} contains a formula whose refs would shift"
                )

        template_spans = sx.merges_within_row(self.merges, op.row_template_row)
        template_styles = self.row_styles(op.row_template_row)

        shifted: dict[int, sx.Row] = {}
        for number, row in self.rows.items():
            if number >= op.at_row:
                moved = sx.shift_row(row, count)
                shifted[moved.number] = moved
            else:
                shifted[number] = row
        self.rows = shifted
        self.merges = sx.shift_merges(self.merges, op.at_row, count)

        placed: list[tuple[int, dict[str, Any]]] = []
        for index, payload in enumerate(op.rows):
            number = op.at_row + index
            new_row = sx.blank_row(sx.retarget_row(template, number))
            for column, style in op.column_styles.items():
                cell = self._cell_in(new_row, f"{column}{number}")
                if cell is not None:
                    new_row = sx.Row(
                        number=number,
                        text=sx.set_cell(new_row.text, sx.restyle_cell(cell, style)),
                    )
            for column, value in payload.items():
                ref = f"{column}{number}"
                value_type = op.value_types.get(column) or infer_value_type(value)
                style = op.column_styles.get(column) or template_styles.get(column)
                self._check_date_style(ref, value_type, style)
                new_row = sx.Row(
                    number=number,
                    text=sx.set_cell(
                        new_row.text,
                        build_cell(ref, value, value_type, style, date1904=self.date1904),
                    ),
                )
                self.max_col_touched = max(
                    self.max_col_touched, column_index_from_string(column)
                )
            self.rows[number] = sx.set_row_spans(new_row)
            for min_col, max_col in template_spans:
                self.merges.append(sx.merge_ref(min_col, max_col, number))
            placed.append((number, payload))

        self.rows_inserted += count
        return placed

    def _cell_in(self, row: sx.Row, ref: str) -> str | None:
        for cell in sx.read_cells(row.text):
            if sx.cell_ref(cell) == ref:
                return cell
        return None

    # -- render --------------------------------------------------------

    def render(self) -> str:
        xml = sx.write_rows(self.xml, list(self.rows.values()))
        # Re-serialising an unchanged <mergeCells> would reorder it for no
        # reason; the fewer bytes that differ from what Excel wrote, the less
        # there is to explain if a file ever does misbehave.
        if self.merges != self.original_merges:
            xml = sx.write_merges(xml, self.merges)
        max_row = max(self.rows) if self.rows else 0
        return sx.extend_dimension(xml, max_row, self.max_col_touched or None)

    @property
    def merge_delta(self) -> int:
        return len(self.merges) - self.merges_before


# ── spec compilation ──────────────────────────────────────────────────────


def _expand(op: ReplaceRegionRowsOp, existing_rows: int) -> list[Operation]:
    """Decompose `replace_region_rows` into primitive ops.

    Fill the rows the block already has; grow it if there are more rows than
    fit; blank the leftovers if there are fewer. Shrinking blanks rather than
    deletes -- the plan's op set has no delete, and an empty styled row is a
    recoverable state where a removed one isn't.
    """
    min_col, min_row, max_col, max_row = sx.bounds(op.ref)
    template_row = op.row_template_row or min_row
    expanded: list[Operation] = []

    for index, payload in enumerate(op.rows[:existing_rows]):
        row_number = min_row + index
        for column, value in payload.items():
            expanded.append(
                SetValueOp(
                    sheet_file=op.sheet_file,
                    ref=f"{column}{row_number}",
                    value=value,
                    value_type=op.value_types.get(column) or infer_value_type(value),
                    style=op.column_styles.get(column),
                )
            )

    if len(op.rows) > existing_rows:
        expanded.append(
            CloneTableRowsOp(
                sheet_file=op.sheet_file,
                at_row=max_row + 1,
                rows=op.rows[existing_rows:],
                row_template_row=template_row,
                column_styles=op.column_styles,
                value_types=op.value_types,
            )
        )
    elif len(op.rows) < existing_rows:
        first_unused = min_row + len(op.rows)
        expanded.append(
            ClearRegionOp(
                sheet_file=op.sheet_file,
                ref=(
                    f"{get_column_letter(min_col)}{first_unused}:"
                    f"{get_column_letter(max_col)}{max_row}"
                ),
            )
        )
    return expanded


def _date1904(archive: zipfile.ZipFile) -> bool:
    workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
    return 'date1904="1"' in workbook_xml or 'date1904="true"' in workbook_xml


def compile_change_spec(
    original_bytes: bytes,
    operations: list[Operation],
    *,
    protected: "list[ProtectedRange] | None" = None,
) -> CompileResult:
    """Apply a Change Spec and report what it did.

    `protected` ranges (from the region schema) are enforced alongside any
    `ProtectOp` in the spec itself: the pre-flight runs before a single byte is
    touched, so a spec that would write the logo fails without producing a
    file at all.
    """
    from app.compiler.protect import preflight  # local import: protect imports ops

    with zipfile.ZipFile(io.BytesIO(original_bytes)) as source:
        date1904 = _date1904(source)
        formats = read_style_formats(source.read("xl/styles.xml").decode("utf-8"))
        editors: dict[str, _SheetEditor] = {}

        def editor_for(sheet_file: str) -> _SheetEditor:
            if sheet_file not in editors:
                editors[sheet_file] = _SheetEditor(
                    source.read(sheet_file).decode("utf-8"),
                    date1904=date1904,
                    formats=formats,
                )
            return editors[sheet_file]

        # `replace_region_rows` needs to know how many rows the block has
        # before it can decide between growing and blanking.
        flat: list[Operation] = []
        for op in operations:
            if isinstance(op, ReplaceRegionRowsOp):
                _, min_row, _, max_row = sx.bounds(op.ref)
                flat.extend(_expand(op, max_row - min_row + 1))
            else:
                flat.append(op)

        preflight(flat, protected or [])

        value_ops = [op for op in flat if isinstance(op, SetValueOp)]
        clear_ops = [op for op in flat if isinstance(op, ClearRegionOp)]
        insert_ops = [op for op in flat if isinstance(op, CloneTableRowsOp)]

        for op in clear_ops:
            editor_for(op.sheet_file).clear(op.ref)
        for op in value_ops:
            editor_for(op.sheet_file).set_value(op)

        # Bottom-up, so an insertion never invalidates the row number of one
        # still to be applied. This is what makes original-document
        # coordinates valid across a whole spec.
        placements: list[tuple[CloneTableRowsOp, list[tuple[int, dict[str, Any]]]]] = []
        for op in sorted(insert_ops, key=lambda op: op.at_row, reverse=True):
            placements.append((op, editor_for(op.sheet_file).insert_rows(op)))

        assertions = _assertions(value_ops, placements, insert_ops)

        out_buffer = io.BytesIO()
        with zipfile.ZipFile(out_buffer, "w", zipfile.ZIP_DEFLATED) as destination:
            rendered = {name: editor.render() for name, editor in editors.items()}
            for item in source.infolist():
                if item.filename in rendered:
                    destination.writestr(item, rendered[item.filename].encode("utf-8"))
                else:
                    destination.writestr(item, source.read(item.filename))

    return CompileResult(
        data=out_buffer.getvalue(),
        rows_inserted=sum(editor.rows_inserted for editor in editors.values()),
        merge_delta=sum(editor.merge_delta for editor in editors.values()),
        operations_applied=len(value_ops) + len(clear_ops) + len(insert_ops),
        assertions=assertions,
    )


def _shift_for(row_number: int, insert_ops: list[CloneTableRowsOp], exclude: int | None) -> int:
    """How far an original-coordinate row moves once all insertions land."""
    return sum(
        len(op.rows)
        for op in insert_ops
        if op.at_row <= row_number and op.at_row != exclude
    )


def _assertions(
    value_ops: list[SetValueOp],
    placements: list[tuple[CloneTableRowsOp, list[tuple[int, dict[str, Any]]]]],
    insert_ops: list[CloneTableRowsOp],
) -> list[Assertion]:
    assertions: list[Assertion] = []
    for op in value_ops:
        if op.value is None or op.value == "":
            continue
        column, row_number = sx.split_ref(op.ref)
        final = row_number + _shift_for(row_number, insert_ops, exclude=None)
        assertions.append(
            Assertion(op.sheet_file, f"{column}{final}", op.value, op.value_type)
        )
    for op, placed in placements:
        for row_number, payload in placed:
            final = row_number + _shift_for(row_number, insert_ops, exclude=op.at_row)
            for column, value in payload.items():
                if value is None or value == "":
                    continue
                assertions.append(
                    Assertion(
                        op.sheet_file,
                        f"{column}{final}",
                        value,
                        op.value_types.get(column) or infer_value_type(value),
                    )
                )
    return assertions


def apply_change_spec(original_bytes: bytes, operations: list[Operation]) -> bytes:
    """Compile a spec and return just the bytes."""
    return compile_change_spec(original_bytes, operations).data
