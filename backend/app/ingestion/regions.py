"""Deterministic region detector for xlsx templates.

These documents have no placeholder markers -- no `{{field}}`, no content
controls. What they have is *shape*: a bold grey label with an empty merged
cell beside it, a styled header row with nothing under it, a serial column
running down the side of a prose column. This module reads that shape and
nothing else. No LLM, no guessing at meaning.

Two rules make it safe to build a writer on top of:

1. **Anything unclassified becomes `immutable`.** An unwritten region is a
   visible gap in the output; a wrongly-written one is silent corruption of
   somebody's form. Only shapes we positively recognise are ever marked
   editable.
2. **Scans are bounded by real content, never by `max_row`.** BFL's
   `Business Function List` sheet reports `max_row=1048575` /
   `max_col=16384` from whole-grid styling while holding ~30 real rows.

Merged ranges, formulas and rich-value anchors are read straight from the
package XML rather than through openpyxl, whose read-only worksheets don't
expose `merged_cells` at all and whose `data_only=True` mode hides formulas
behind their cached values.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

DETECTOR_VERSION = "structural/1.0.0"

# ── Scan bounds ────────────────────────────────────────────────────────────
# Generous enough for every sheet in the corpus (largest real extent is 604
# rows), tight enough that a whole-grid-styled sheet can't run away with the
# worker.
MAX_SCAN_ROWS = 5_000
MAX_SCAN_COLS = 64
# BFL's Cover sheet holds its only content at rows 16-18 inside a declared
# A1:Z1000 extent, so the streak has to tolerate a long leading gap.
BLANK_ROW_STREAK = 50

# ── Detection thresholds ──────────────────────────────────────────────────
TITLE_MAX_ROW = 3  # a merged banner this high up is the form's own title
TITLE_WIDTH_RATIO = 0.5  # ... if it spans at least half the used width
MIN_HEADER_CELLS = 3
MAX_LABEL_LEN = 60  # longer than this is prose, not a label
# Column headers run longer than field labels -- BFL's
# `Frequency(Ondemand,Realtime,Daily,Weekly,Monthly,Yearly,Others)` is 63
# characters and is unmistakably a header.
MAX_HEADER_CELL_LEN = 120
MIN_PROSE_LEN = 40  # shorter than this isn't a narrative cell
MIN_SERIAL_RUN = 3
MAX_SECTION_LABEL_LOOKUP = 2  # rows to walk up looking for a block's caption

# YMSLI's templates park a literal '*' in cells to mark print extents. Counted
# as content it turns a table_loop into a table_data with a dozen phantom rows.
SENTINEL_VALUES = frozenset({"*"})

_SHEET_TAG_RE = re.compile(r"<sheet\b[^>]*?>")
_NAME_ATTR_RE = re.compile(r'\bname="([^"]*)"')
_RID_ATTR_RE = re.compile(r'\br:id="([^"]*)"')
_REL_TAG_RE = re.compile(r"<Relationship\b[^>]*?>")
_ID_ATTR_RE = re.compile(r'\bId="([^"]+)"')
_TARGET_ATTR_RE = re.compile(r'\bTarget="([^"]+)"')
_MERGE_REF_RE = re.compile(r'<mergeCell\s+ref="([^"]+)"')
_CELL_OPEN_RE = re.compile(r"<c\s([^>]*?)(/>|>)")
_CELL_REF_ATTR_RE = re.compile(r'\br="([A-Z]+\d+)"')
_VM_ATTR_RE = re.compile(r'\bvm="\d+"')
_DRAWING_RID_RE = re.compile(r'<drawing\s+r:id="([^"]+)"')
_ANCHOR_BLOCK_RE = re.compile(
    r"<xdr:(twoCellAnchor|oneCellAnchor|absoluteAnchor)\b.*?</xdr:\1>", re.DOTALL
)
_ANCHOR_FROM_RE = re.compile(
    r"<xdr:from>\s*<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>", re.DOTALL
)
_NON_ALNUM_RE = re.compile(r"[^a-z0-9]+")

# kind -> (editable, protected). `protected` is what the compiler pre-flight
# refuses to write through; `editable` is what the generation agent may offer.
_WRITABILITY: dict[str, tuple[bool, bool]] = {
    "field": (True, False),
    "field_empty": (True, False),
    "table_loop": (True, False),
    "table_data": (True, False),
    "repeating_block": (True, False),
    "title": (False, False),
    "history": (False, False),  # system-written on publish, not user-written
    "branding": (False, True),
    "derived": (False, True),
    "instruction": (False, True),
    "immutable": (False, True),
}


@dataclass
class Locator:
    """Where a region lives. Mirrors `document_regions.locator` JSONB."""

    sheet: str
    ref: str  # "B5", "A27:G27", "A11:G24"
    sheet_file: str | None = None  # "xl/worksheets/sheet1.xml", for the compiler


@dataclass
class Region:
    region_id: str
    kind: str
    label: str | None
    locator: Locator
    value_type: str  # string | number | date | mixed | none
    current_content: str | None
    editable: bool
    protected: bool
    required: bool
    detection_method: str
    ordinal: int
    notes: dict[str, Any] = field(default_factory=dict)


@dataclass
class SheetSchema:
    name: str
    sheet_file: str
    role: str | None  # cover | history | instruction | None
    used_range: str | None
    regions: list[Region]
    notes: dict[str, Any] = field(default_factory=dict)


@dataclass
class DocumentSchema:
    detector_version: str
    format: str
    sheets: list[SheetSchema]

    @property
    def regions(self) -> list[Region]:
        return [region for sheet in self.sheets for region in sheet.regions]

    def counts_by_kind(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        for region in self.regions:
            counts[region.kind] = counts.get(region.kind, 0) + 1
        return dict(sorted(counts.items()))

    def writable(self) -> list[Region]:
        return [region for region in self.regions if region.editable]


# ── Package-level facts, read from the zip ────────────────────────────────


@dataclass
class _SheetParts:
    name: str
    sheet_file: str
    merges: list[tuple[int, int, int, int]]  # (min_col, min_row, max_col, max_row)
    formula_refs: set[str]
    rich_value_refs: set[str]
    drawing_anchors: dict[tuple[int, int], str]  # (row, col) 1-based -> picture | shape


def _scan_cells(sheet_xml: str) -> tuple[set[str], set[str]]:
    """Return (formula refs, rich-value refs) from one worksheet's XML.

    Deliberately a linear scan rather than a `<c>...</c>` regex: these sheets
    carry tens of thousands of *self-closing* `<c/>` elements from whole-grid
    styling, and any pattern that looks past the tag for a closing `</c>`
    backtracks across the whole 500 KB part on every one of them.
    """
    formula_refs: set[str] = set()
    rich_value_refs: set[str] = set()

    for match in _CELL_OPEN_RE.finditer(sheet_xml):
        attrs, terminator = match.group(1), match.group(2)
        ref_match = _CELL_REF_ATTR_RE.search(attrs)
        if ref_match is None:
            continue
        ref = ref_match.group(1)
        if _VM_ATTR_RE.search(attrs):
            rich_value_refs.add(ref)
        if terminator == "/>":
            continue
        close = sheet_xml.find("</c>", match.end())
        if close != -1 and "<f" in sheet_xml[match.end() : close]:
            formula_refs.add(ref)

    return formula_refs, rich_value_refs


def _rels_for(archive: zipfile.ZipFile, part: str) -> dict[str, str]:
    prefix, _, name = part.rpartition("/")
    rels_part = f"{prefix}/_rels/{name}.rels"
    if rels_part not in archive.namelist():
        return {}
    xml = archive.read(rels_part).decode("utf-8")
    # Attribute order isn't fixed: Excel writes Id before Target, openpyxl
    # doesn't, so each attribute is pulled out of the tag separately.
    rels: dict[str, str] = {}
    for tag in _REL_TAG_RE.findall(xml):
        rel_id = _ID_ATTR_RE.search(tag)
        target = _TARGET_ATTR_RE.search(tag)
        if rel_id and target:
            rels[rel_id.group(1)] = target.group(1)
    return rels


def _resolve(base_part: str, target: str) -> str:
    """Resolve a relationship target against the part that declared it."""
    if target.startswith("/"):
        return target.lstrip("/")
    prefix = base_part.rpartition("/")[0]
    parts: list[str] = []
    for segment in f"{prefix}/{target}".split("/"):
        if segment == "..":
            if parts:
                parts.pop()
        elif segment not in ("", "."):
            parts.append(segment)
    return "/".join(parts)


def _read_sheet_parts(data: bytes) -> list[_SheetParts]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = set(archive.namelist())
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        workbook_rels = _rels_for(archive, "xl/workbook.xml")

        sheets: list[_SheetParts] = []
        for tag in _SHEET_TAG_RE.findall(workbook_xml):
            name_match = _NAME_ATTR_RE.search(tag)
            rid_match = _RID_ATTR_RE.search(tag)
            if name_match is None or rid_match is None:
                continue
            target = workbook_rels.get(rid_match.group(1))
            if target is None:
                continue
            sheet_file = _resolve("xl/workbook.xml", target)
            if sheet_file not in names:
                continue

            sheet_xml = archive.read(sheet_file).decode("utf-8")
            merges = []
            for ref in _MERGE_REF_RE.findall(sheet_xml):
                bounds = range_boundaries(ref)
                merges.append((bounds[0], bounds[1], bounds[2], bounds[3]))

            anchors: dict[tuple[int, int], str] = {}
            sheet_rels = _rels_for(archive, sheet_file)
            for rid in _DRAWING_RID_RE.findall(sheet_xml):
                target = sheet_rels.get(rid)
                if target is None:
                    continue
                drawing_part = _resolve(sheet_file, target)
                if drawing_part not in names:
                    continue
                drawing_xml = archive.read(drawing_part).decode("utf-8")
                for block_match in _ANCHOR_BLOCK_RE.finditer(drawing_xml):
                    block = block_match.group(0)
                    origin = _ANCHOR_FROM_RE.search(block)
                    if origin is None:  # absoluteAnchor: positioned in EMUs, no cell
                        continue
                    cell_key = (int(origin.group(2)) + 1, int(origin.group(1)) + 1)
                    anchors[cell_key] = "picture" if "<xdr:pic" in block else "shape"

            formula_refs, rich_value_refs = _scan_cells(sheet_xml)
            sheets.append(
                _SheetParts(
                    name=name_match.group(1),
                    sheet_file=sheet_file,
                    merges=merges,
                    formula_refs=formula_refs,
                    rich_value_refs=rich_value_refs,
                    drawing_anchors=anchors,
                )
            )
    return sheets


# ── Grid ──────────────────────────────────────────────────────────────────


@dataclass
class _Cell:
    row: int
    col: int
    value: Any
    text: str
    bold: bool
    filled: bool
    style_key: str
    number_format: str


def _style_of(cell: Any) -> tuple[bool, bool, str]:
    try:
        bold = bool(cell.font.bold)
    except Exception:  # pragma: no cover - malformed style records
        bold = False
    pattern = ""
    colour = ""
    try:
        pattern = cell.fill.patternType or ""
        if pattern:
            fg = cell.fill.fgColor
            colour = f"{fg.type}:{getattr(fg, fg.type, '')}"
    except Exception:  # pragma: no cover
        pattern, colour = "", ""
    return bold, bool(pattern), f"{bold}|{pattern}|{colour}"


def _read_grid(worksheet: Any) -> tuple[dict[tuple[int, int], _Cell], int, int]:
    """Return {(row, col): cell} for non-empty cells, plus last row/col used.

    Bounded by BLANK_ROW_STREAK consecutive empty rows and by MAX_SCAN_ROWS /
    MAX_SCAN_COLS, so a sheet that claims a million rows can't stall ingestion.
    """
    grid: dict[tuple[int, int], _Cell] = {}
    last_row = 0
    last_col = 0
    blank_streak = 0

    for row_index, row in enumerate(worksheet.iter_rows(max_col=MAX_SCAN_COLS), start=1):
        if row_index > MAX_SCAN_ROWS:
            break
        row_had_content = False
        for col_index, cell in enumerate(row, start=1):
            value = getattr(cell, "value", None)
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            bold, filled, style_key = _style_of(cell)
            grid[(row_index, col_index)] = _Cell(
                row=row_index,
                col=col_index,
                value=value,
                text=str(value).strip(),
                bold=bold,
                filled=filled,
                style_key=style_key,
                number_format=getattr(cell, "number_format", "General") or "General",
            )
            row_had_content = True
            last_col = max(last_col, col_index)
        if row_had_content:
            last_row = row_index
            blank_streak = 0
        else:
            blank_streak += 1
            if blank_streak >= BLANK_ROW_STREAK:
                break

    return grid, last_row, last_col


# ── Small helpers ─────────────────────────────────────────────────────────


def _ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _range_ref(min_col: int, min_row: int, max_col: int, max_row: int) -> str:
    start = _ref(min_row, min_col)
    if (min_col, min_row) == (max_col, max_row):
        return start
    return f"{start}:{_ref(max_row, max_col)}"


def _slug(text: str, fallback: str) -> str:
    slug = _NON_ALNUM_RE.sub("_", text.lower()).strip("_")
    return slug[:48] or fallback


def _normalise_sheet_name(name: str) -> str:
    """`History` and `History.` are the same sheet role; match fuzzily."""
    return _NON_ALNUM_RE.sub("", name.lower())


def _sheet_role(name: str) -> str | None:
    normalised = _normalise_sheet_name(name)
    for role in ("history", "instruction", "cover"):
        if role in normalised:
            return role
    return None


def _is_sentinel(cell: _Cell) -> bool:
    return cell.text in SENTINEL_VALUES


def _is_serial(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _value_type(cell: _Cell | None) -> str:
    if cell is None:
        return "none"
    if isinstance(cell.value, (datetime, date)):
        return "date"
    if isinstance(cell.value, bool):
        return "string"
    if isinstance(cell.value, (int, float)):
        # A serial date only looks numeric; the number format is the giveaway.
        fmt = cell.number_format.lower()
        if any(token in fmt for token in ("yy", "mmm", "d-m", "dd/")):
            return "date"
        return "number"
    return "string"


def _looks_like_label(cell: _Cell) -> bool:
    if not isinstance(cell.value, str):
        return False
    if len(cell.text) > MAX_LABEL_LEN or _is_sentinel(cell):
        return False
    return cell.bold or cell.filled or cell.text.endswith(":")


# ── Per-sheet detection ───────────────────────────────────────────────────


class _SheetDetector:
    def __init__(self, parts: _SheetParts, worksheet: Any, used_ids: set[str]) -> None:
        self.parts = parts
        self.role = _sheet_role(parts.name)
        self.grid, self.last_row, self.last_col = _read_grid(worksheet)
        self.claimed: set[tuple[int, int]] = set()
        self.regions: list[Region] = []
        self.sentinel_cells = 0
        # Shared across sheets: `document_regions` is UNIQUE (version_id,
        # region_id), and `History` / `History.` both slug to 'history'.
        self._used_ids = used_ids
        # header row -> (min_col, max_col), filled in by _detect_tables
        self._header_rows: dict[int, tuple[int, int]] = {}

        # (row, col) -> the merged block it belongs to
        self.merge_of: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for min_col, min_row, max_col, max_row in parts.merges:
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    self.merge_of[(row, col)] = (min_col, min_row, max_col, max_row)

    # -- claiming ---------------------------------------------------------

    def _claim(self, min_col: int, min_row: int, max_col: int, max_row: int) -> None:
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                self.claimed.add((row, col))

    def _is_claimed(self, row: int, col: int) -> bool:
        return (row, col) in self.claimed

    def _extent(self, row: int, col: int) -> tuple[int, int, int, int]:
        """The merged block at (row, col), or the single cell itself."""
        return self.merge_of.get((row, col), (col, row, col, row))

    def _cell(self, row: int, col: int) -> _Cell | None:
        """The cell at (row, col), following a merge back to its anchor."""
        cell = self.grid.get((row, col))
        if cell is not None:
            return cell
        merge = self.merge_of.get((row, col))
        if merge is not None:
            return self.grid.get((merge[1], merge[0]))
        return None

    def _row_cells(self, row: int, min_col: int = 1, max_col: int | None = None) -> list[_Cell]:
        upper = self.last_col if max_col is None else max_col
        return [
            self.grid[(row, col)] for col in range(min_col, upper + 1) if (row, col) in self.grid
        ]

    # -- emission ---------------------------------------------------------

    def _emit(
        self,
        kind: str,
        locator_ref: str,
        *,
        label: str | None = None,
        value_type: str = "none",
        current_content: str | None = None,
        id_hint: str | None = None,
        notes: dict[str, Any] | None = None,
    ) -> Region:
        if self.role == "instruction":
            # The Instruction sheet documents the template's own fill rules.
            # Nothing in it is ever a write target.
            kind = "instruction"
        editable, protected = _WRITABILITY[kind]
        sheet_slug = _slug(self.parts.name, "sheet")
        local_slug = _slug(id_hint or label or locator_ref, kind)
        if local_slug == sheet_slug:  # 'business_function_list.business_function_list'
            local_slug = kind
        base = f"{sheet_slug}.{local_slug}"
        region_id = base
        suffix = 2
        while region_id in self._used_ids:
            region_id = f"{base}_{suffix}"
            suffix += 1
        self._used_ids.add(region_id)

        region = Region(
            region_id=region_id,
            kind=kind,
            label=label,
            locator=Locator(
                sheet=self.parts.name, ref=locator_ref, sheet_file=self.parts.sheet_file
            ),
            value_type=value_type,
            current_content=current_content,
            editable=editable,
            protected=protected,
            required=False,  # the LLM pass owns requiredness; structure can't tell
            detection_method="structural",
            ordinal=len(self.regions),
            notes=notes or {},
        )
        self.regions.append(region)
        return region

    # -- passes -----------------------------------------------------------

    def detect(self) -> SheetSchema:
        if self.last_row:
            # Order is a priority ladder: the never-writable shapes claim their
            # cells first, then the block shapes, then single fields, and
            # whatever is left over defaults to immutable. Tables come before
            # titles because BFL parks a merged data cell at `C3:O3` -- inside
            # the table, but shallow enough to pass for a banner.
            self._detect_anchors()
            self._detect_derived()
            self._detect_tables()
            self._detect_repeating_blocks()
            self._detect_titles()
            self._detect_fields()
            self._detect_immutable()

        # Detection runs in priority order (branding before tables before
        # fields); readers and the question planner want document order.
        self.regions.sort(key=lambda region: range_boundaries(region.locator.ref)[1::-1])
        for ordinal, region in enumerate(self.regions):
            region.ordinal = ordinal

        used = (
            _range_ref(1, 1, self.last_col, self.last_row)
            if self.last_row and self.last_col
            else None
        )
        return SheetSchema(
            name=self.parts.name,
            sheet_file=self.parts.sheet_file,
            role=self.role,
            used_range=used,
            regions=self.regions,
            notes={"sentinel_cells": self.sentinel_cells} if self.sentinel_cells else {},
        )

    def _detect_anchors(self) -> None:
        """In-cell rich values (the MOM logo) and drawing anchors. Never written.

        An anchored *picture* is branding. An anchored *shape* is a text box or
        connector -- also untouchable, but calling it branding would be a lie,
        so it takes the fail-safe `immutable`.
        """
        anchors: list[tuple[int, int, str, str]] = []
        for ref in sorted(self.parts.rich_value_refs):
            bounds = range_boundaries(ref)
            anchors.append((bounds[1], bounds[0], "branding", "rich_value"))
        for (row, col), anchor_kind in sorted(self.parts.drawing_anchors.items()):
            kind = "branding" if anchor_kind == "picture" else "immutable"
            anchors.append((row, col, kind, f"drawing_{anchor_kind}"))

        for row, col, kind, source in anchors:
            if self._is_claimed(row, col):
                continue
            min_col, min_row, max_col, max_row = self._extent(row, col)
            self._claim(min_col, min_row, max_col, max_row)
            cell = self._cell(row, col)
            self._emit(
                kind,
                _range_ref(min_col, min_row, max_col, max_row),
                label=None,
                current_content=cell.text if cell else None,
                id_hint=f"{source}_{_ref(row, col).lower()}",
                notes={"anchor": source},
            )

    def _detect_derived(self) -> None:
        for ref in sorted(self.parts.formula_refs):
            bounds = range_boundaries(ref)
            col, row = bounds[0], bounds[1]
            if self._is_claimed(row, col):
                continue
            cell = self._cell(row, col)
            self._claim(col, row, col, row)
            self._emit(
                "derived",
                ref,
                value_type=_value_type(cell),
                current_content=cell.text if cell else None,
                id_hint=f"derived_{ref.lower()}",
            )

    def _detect_titles(self) -> None:
        min_width = max(2, int(self.last_col * TITLE_WIDTH_RATIO))
        for min_col, min_row, max_col, max_row in sorted(
            self.parts.merges, key=lambda m: (m[1], m[0])
        ):
            if min_row > TITLE_MAX_ROW or (max_col - min_col + 1) < min_width:
                continue
            cell = self.grid.get((min_row, min_col))
            if cell is None or self._is_claimed(min_row, min_col):
                continue
            self._claim(min_col, min_row, max_col, max_row)
            self._emit(
                "title",
                _range_ref(min_col, min_row, max_col, max_row),
                label=cell.text,
                value_type="string",
                current_content=cell.text,
                id_hint=f"title_{_slug(cell.text, _ref(min_row, min_col).lower())}",
            )

    # -- tables -----------------------------------------------------------

    def _header_run(self, row: int) -> tuple[int, int, list[_Cell]] | None:
        """Is `row` a table header? Return (min_col, max_col, cells) if so.

        A header row and a label/value row look alike until you check two
        things. A header's cells share one style (MOM row 27 is seven bold
        cells); a label/value row alternates (`A6` bold grey "Prepared By",
        `B6` plain "YMSLI"). And a header's cells aren't merged, while this
        corpus builds label/value rows *out of* merges (`B6:C6`, `E6:G6`).
        """
        cells = [
            cell
            for cell in self._row_cells(row)
            if not self._is_claimed(cell.row, cell.col) and not _is_sentinel(cell)
        ]
        if len(cells) < MIN_HEADER_CELLS:
            return None
        if not all(
            isinstance(cell.value, str) and len(cell.text) <= MAX_HEADER_CELL_LEN
            for cell in cells
        ):
            return None
        if len({cell.style_key for cell in cells}) != 1:
            return None
        if any(self.merge_of.get((row, cell.col)) for cell in cells):
            return None

        min_col = min(cell.col for cell in cells)
        max_col = max(cell.col for cell in cells)
        # Tolerate the odd empty header cell (BFL's column D) but not a row of
        # scattered labels.
        if (max_col - min_col + 1) > 2 * len(cells):
            return None
        return min_col, max_col, cells

    def _count_data_rows(self, header_row: int, min_col: int, max_col: int) -> int:
        count = 0
        row = header_row + 1
        while row <= self.last_row:
            populated = [
                cell
                for cell in self._row_cells(row, min_col, max_col)
                if not _is_sentinel(cell)
            ]
            if not populated:
                break
            count += 1
            row += 1
        return count

    def _section_label_above(self, row: int, min_col: int, max_col: int) -> _Cell | None:
        """The caption sitting directly above a block, e.g. MOM `A26`."""
        if row <= 1:
            return None
        candidates = [
            cell
            for cell in self._row_cells(row - 1, max(1, min_col - 1), max_col)
            if not self._is_claimed(cell.row, cell.col) and _looks_like_label(cell)
        ]
        if len(candidates) != 1:
            return None
        return candidates[0]

    def _detect_tables(self) -> None:
        row = 1
        while row <= self.last_row:
            run = self._header_run(row)
            if run is None:
                row += 1
                continue
            min_col, max_col, cells = run
            self._header_rows[row] = (min_col, max_col)
            data_rows = self._count_data_rows(row, min_col, max_col)
            kind = "table_data" if data_rows else "table_loop"
            if self.role == "history":
                # The revision block is ours to append to on publish, not the
                # generation agent's to fill in.
                kind = "history"

            columns = [
                {"ref": _ref(row, cell.col), "column": get_column_letter(cell.col), "name": cell.text}
                for cell in cells
            ]
            header_ref = _range_ref(min_col, row, max_col, row)
            last_data_row = row + data_rows
            caption = self._section_label_above(row, min_col, max_col)

            self._claim(min_col, row, max_col, last_data_row)
            if caption is not None:
                extent = self._extent(caption.row, caption.col)
                self._claim(*extent)

            notes: dict[str, Any] = {
                "header_ref": header_ref,
                "columns": columns,
                "row_count": data_rows,
                "first_data_row": row + 1,
            }
            if kind == "table_loop":
                notes["insert_at_row"] = row + 1
            if kind == "history":
                notes["system_written"] = True
            # A table with no caption above it is named by its sheet -- POC's
            # `YMVN POC list` is the whole point of that sheet.
            label = caption.text if caption else self.parts.name
            self._emit(
                kind,
                _range_ref(min_col, row, max_col, last_data_row),
                label=label,
                value_type="mixed",
                id_hint=label,
                notes=notes,
            )
            row = last_data_row + 1

    # -- repeating blocks -------------------------------------------------

    def _detect_repeating_blocks(self) -> None:
        for col in range(1, self.last_col + 1):
            row = 1
            while row <= self.last_row:
                run = self._serial_run(col, row)
                if run is None:
                    row += 1
                    continue
                start_row, end_row, serials = run
                content_col = self._prose_column(col, start_row, end_row)
                if content_col is None:
                    row = end_row + 1
                    continue

                block_max_col = content_col
                merge = self.merge_of.get((start_row, content_col))
                if merge is not None:
                    block_max_col = merge[2]

                self._claim(col, start_row, block_max_col, end_row)
                label_cell = self._section_label_above(start_row, col, block_max_col)
                column_header = None
                if label_cell is not None:
                    self._claim(*self._extent(label_cell.row, label_cell.col))
                    column_header = label_cell.text
                    outer = self._section_label_above(label_cell.row, col, block_max_col)
                    if outer is not None:
                        self._claim(*self._extent(outer.row, outer.col))
                        label_cell = outer

                gaps = [
                    [serials[i], serials[i + 1]]
                    for i in range(len(serials) - 1)
                    if serials[i + 1] != serials[i] + 1
                ]
                self._emit(
                    "repeating_block",
                    _range_ref(col, start_row, block_max_col, end_row),
                    label=label_cell.text if label_cell else None,
                    value_type="string",
                    id_hint=(label_cell.text if label_cell else None) or f"block_row{start_row}",
                    notes={
                        "serial_column": get_column_letter(col),
                        "content_column": get_column_letter(content_col),
                        "column_header": column_header,
                        "row_count": end_row - start_row + 1,
                        "serial_values": serials,
                        # Appendix B: the MOM's Sl.No. jumps 13 -> 16. Surface
                        # it; renumbering silently would rewrite the user's doc.
                        "serial_gaps": gaps,
                        "insert_at_row": end_row + 1,
                    },
                )
                row = end_row + 1

    def _serial_run(self, col: int, start_row: int) -> tuple[int, int, list[int]] | None:
        serials: list[int] = []
        row = start_row
        while row <= self.last_row:
            if self._is_claimed(row, col):
                break
            cell = self.grid.get((row, col))
            if cell is None:
                break
            serial = _is_serial(cell.value)
            if serial is None or (serials and serial <= serials[-1]):
                break
            serials.append(serial)
            row += 1
        if len(serials) < MIN_SERIAL_RUN:
            return None
        return start_row, start_row + len(serials) - 1, serials

    def _prose_column(self, serial_col: int, start_row: int, end_row: int) -> int | None:
        """The narrative column beside a serial run, if there is one."""
        rows = end_row - start_row + 1
        for col in range(serial_col + 1, min(serial_col + 3, self.last_col) + 1):
            prose_rows = sum(
                1
                for row in range(start_row, end_row + 1)
                if (cell := self._cell(row, col)) is not None
                and isinstance(cell.value, str)
                and len(cell.text) >= MIN_PROSE_LEN
            )
            if prose_rows * 2 >= rows:
                return col
        return None

    # -- fields -----------------------------------------------------------

    def _detect_fields(self) -> None:
        for (row, col), cell in sorted(self.grid.items()):
            if self._is_claimed(row, col) or not _looks_like_label(cell):
                continue
            label_extent = self._extent(row, col)
            label_width = label_extent[2] - label_extent[0] + 1
            if label_width >= max(2, int(self.last_col * TITLE_WIDTH_RATIO)):
                # A full-width merged caption below the title rows is section
                # text, not a label with a value beside it. Leave it immutable.
                continue
            if self._captions_a_table(label_extent):
                continue

            target = self._field_target(label_extent)
            if target is None:
                continue
            target_extent, target_cell = target
            kind = "field" if target_cell is not None else "field_empty"
            self._claim(*label_extent)
            self._claim(*target_extent)
            self._emit(
                kind,
                _range_ref(*target_extent),
                label=cell.text,
                value_type=_value_type(target_cell),
                current_content=target_cell.text if target_cell else None,
                id_hint=cell.text,
                notes={"label_ref": _range_ref(*label_extent)},
            )

    def _captions_a_table(self, label_extent: tuple[int, int, int, int]) -> bool:
        """Is this label a group caption sitting on top of a table header row?

        BFL row 1 carries two of them -- `Template Definition Business` and
        `Template Definition (System)` over the header at row 2. Read as field
        labels they'd offer the banner's empty right-hand cells as write
        targets, which would drop text into the middle of a group heading.
        """
        min_col, _, max_col, max_row = label_extent
        span = self._header_rows.get(max_row + 1)
        return span is not None and min_col <= span[1] and max_col >= span[0]

    def _field_target(
        self, label_extent: tuple[int, int, int, int]
    ) -> tuple[tuple[int, int, int, int], _Cell | None] | None:
        """The value slot for a label: to its right, else directly beneath."""
        min_col, min_row, max_col, max_row = label_extent
        for row, col in ((min_row, max_col + 1), (max_row + 1, min_col)):
            if row > self.last_row or col > self.last_col:
                continue
            if self._is_claimed(row, col):
                continue
            extent = self._extent(row, col)
            cell = self._cell(row, col)
            if cell is not None and _is_sentinel(cell):
                continue
            # Styling can't separate label from value here: BFL's Cover sheet
            # sets `Project Name` and `YNS-FO-GT` in the same bold. The layout
            # can -- the value sits in a merged slot (`C16:G16`), the label
            # doesn't. So an unmerged neighbour that reads like a label is
            # taken to be the next label rather than this one's value.
            merged_target = (extent[2] - extent[0] > 0) or (extent[3] - extent[1] > 0)
            if cell is not None and not merged_target and _looks_like_label(cell):
                continue
            return extent, cell
        return None

    def _detect_immutable(self) -> None:
        """Fail-safe default: every remaining populated cell is off limits."""
        for (row, col), cell in sorted(self.grid.items()):
            if self._is_claimed(row, col):
                continue
            if _is_sentinel(cell):
                # Print-extent markers, not content. Counted on the sheet so
                # they're reported rather than silently swallowed.
                self.sentinel_cells += 1
                continue
            extent = self._extent(row, col)
            self._claim(*extent)
            self._emit(
                "immutable",
                _range_ref(*extent),
                label=cell.text if len(cell.text) <= MAX_LABEL_LEN else None,
                value_type=_value_type(cell),
                current_content=cell.text,
                id_hint=f"immutable_{_ref(row, col).lower()}",
            )


def detect_regions_xlsx(data: bytes) -> DocumentSchema:
    """Classify every populated region of an xlsx workbook, structurally.

    `read_only=True` and `data_only=True`: the first keeps a whole-grid-styled
    sheet from exhausting memory, the second means we compare cached values
    rather than formula strings (formulas come from the XML instead).
    """
    sheets_parts = _read_sheet_parts(data)
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        schemas: list[SheetSchema] = []
        used_ids: set[str] = set()
        for parts in sheets_parts:
            if parts.name not in workbook.sheetnames:
                continue
            schemas.append(_SheetDetector(parts, workbook[parts.name], used_ids).detect())
    finally:
        workbook.close()

    return DocumentSchema(
        detector_version=DETECTOR_VERSION, format="xlsx", sheets=schemas
    )
